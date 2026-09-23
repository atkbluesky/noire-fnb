# -*- coding: utf-8 -*-
"""
KIỂM TRA MẪU FILE ĐẦU VÀO
=========================
    python tools/l0_validate.py            kiểm mọi file trong L0_input
    python tools/l0_validate.py S02_bill   kiểm một nguồn

So từng file với `schema` của nguồn trong data_sources.json (khai ở tools/l0_registry.py):
sheet bắt buộc, cột bắt buộc trên cùng một dòng tiêu đề, dấu mốc bảng, file bắt buộc
trong thư mục tháng.

Vì sao cần: iPOS đổi tên cột từ T8/2026 (`Mã hoá đơn→Hoá đơn`, `Thời gian→Ngày`…) và
lane cũ lặng lẽ mất mã hoá đơn + ngày của hai tháng. Chốt này bắt loại lỗi đó NGAY KHI
THẢ FILE, bằng một câu tiếng Việt chỉ rõ thiếu sheet/cột nào.

Chỉ đọc ~30 dòng đầu mỗi sheet; kết quả nhớ theo chữ ký file ở _cache/l0_validate.json
nên file không đổi thì không mở lại.
"""
from __future__ import annotations

import csv
import fnmatch
import io
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from monthly_lib import (  # noqa: E402
    L0_ROOT, ROOT, SOURCES, file_sig, l0_by_month, l0_files, norm, read_html_table,
)

CACHE = os.path.join(ROOT, "_cache", "l0_validate.json")
SCAN_ROWS = 30
VERSION = "1"


def _alts(x):
    return [a for a in (x if isinstance(x, list) else [x])]


def _cell_match(cell, name):
    c, n = norm(cell), norm(name)
    return bool(c) and (c == n or (len(n) > 5 and c.startswith(n)))


def _has_cols(rows, cols):
    """Có một dòng chứa đủ `cols` (mỗi phần tử: tên hoặc danh sách tên thay thế)?
    → (True, None) hoặc (False, [cột thiếu ở dòng khớp nhiều nhất])."""
    best = None
    for r in rows:
        cells = [c for c in r if c is not None and str(c).strip()]
        if not cells:
            continue
        miss = [(" / ".join(_alts(c)) if isinstance(c, list) else c) for c in cols
                if not any(_cell_match(x, a) for x in cells for a in _alts(c))]
        if not miss:
            return True, None
        if best is None or len(miss) < len(best):
            best = miss
    return False, best or [(" / ".join(_alts(c)) if isinstance(c, list) else c) for c in cols]


def _sheet_rows(ws):
    try:
        ws.reset_dimensions()
    except Exception:  # noqa: BLE001
        pass
    out = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i >= SCAN_ROWS:
            break
        out.append([str(c).replace("​", "").strip() if c is not None else None for c in (r or [])])
    return out


def validate_file(sid, path):
    """→ danh sách lỗi (rỗng = đúng mẫu)."""
    sch = SOURCES[sid].get("schema") or {}
    if not sch:
        return []
    ext = os.path.splitext(path)[1].lower()
    if sch.get("only_ext") and ext != sch["only_ext"]:
        return []
    errs = []
    if "csv_cols" in sch:
        try:
            head = next(csv.reader(io.open(path, encoding="utf-8-sig")))
        except Exception as e:  # noqa: BLE001
            return [f"không đọc được CSV: {str(e)[:60]}"]
        miss = [c for c in sch["csv_cols"] if c not in [h.strip() for h in head]]
        return [f"thiếu cột {', '.join(miss)}"] if miss else []
    if "html_any" in sch:
        try:
            header, _ = read_html_table(path)
        except Exception as e:  # noqa: BLE001
            return [f"không đọc được (file OA phải là bản xuất gốc, đừng mở rồi lưu lại): {str(e)[:50]}"]
        if not any(_cell_match(h, n) for h in header for n in sch["html_any"]):
            return [f"không thấy cột nào trong: {', '.join(sch['html_any'])}"]
        return []
    if ext not in (".xlsx", ".xlsm"):
        return []
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        return [f"không mở được file Excel (hỏng hoặc đang chép dở): {str(e)[:60]}"]
    try:
        names = wb.sheetnames
        chosen = []
        for group in sch.get("sheets") or []:
            alts = group if isinstance(group, list) else [group]
            hit = None
            for a in alts:
                if a is None:
                    tot = [n for n in names if norm(n) == "tất cả cửa hàng"]
                    hit = tot[0] if tot else names[0]
                    break
                # `*` = phần đổi theo kỳ trong tên sheet (`03_Q4 Master Plan` ↔ `*Master Plan`)
                m = [n for n in names if (fnmatch.fnmatch(norm(n), norm(a)) if "*" in a else norm(n) == norm(a))]
                if m:
                    hit = m[0]
                    break
            if hit is None:
                errs.append(f"thiếu sheet {' hoặc '.join(repr(a) for a in alts if a)} "
                            f"(file có: {', '.join(names[:6])})")
            else:
                chosen.append(hit)
        target = chosen[0] if chosen else (names[0] if names else None)
        if target and (sch.get("cols") or sch.get("markers") or sch.get("cols_by_sheet")):
            rows = _sheet_rows(wb[target])
            cols = (sch.get("cols_by_sheet") or {}).get(target) or sch.get("cols")
            if cols:
                ok, miss = _has_cols(rows, cols)
                if not ok:
                    errs.append(f"sheet '{target}' thiếu cột: {', '.join(miss)}")
            for mk in sch.get("markers") or []:
                if not any(c and norm(mk) in norm(c) for r in rows for c in r):
                    errs.append(f"sheet '{target}' không thấy mốc '{mk}'")
    finally:
        wb.close()
    return errs


def _load_cache():
    try:
        d = json.load(io.open(CACHE, encoding="utf-8"))
        return d if d.get("v") == VERSION else {"v": VERSION, "files": {}}
    except (OSError, ValueError):
        return {"v": VERSION, "files": {}}


def validate_all(only=None):
    """→ {sid: [(đường_dẫn_tương_đối, [lỗi])]} chỉ gồm file SAI mẫu + thư mục tháng thiếu file."""
    cache = _load_cache()
    out = {}
    for sid, s in SOURCES.items():
        if only and sid not in only:
            continue
        sch = s.get("schema") or {}
        bad = []
        for f in l0_files(sid):
            if sch.get("month_files"):
                continue
            rel = os.path.relpath(f, L0_ROOT).replace("\\", "/")
            sig = f"{file_sig(f)}|{json.dumps(sch, ensure_ascii=False, sort_keys=True)}"
            hit = cache["files"].get(rel)
            if hit and hit["sig"] == sig:
                errs = hit["errs"]
            else:
                errs = validate_file(sid, f)
                cache["files"][rel] = {"sig": sig, "errs": errs}
            if errs:
                bad.append((rel, errs))
        if sch.get("month_files"):
            for m, f in sorted(l0_by_month(sid).items()):
                d = os.path.dirname(f)
                have = {norm(x) for x in os.listdir(d)}
                miss = [x for x in sch["month_files"] if norm(x) not in have]
                if miss:
                    bad.append((os.path.relpath(d, L0_ROOT).replace("\\", "/"),
                                [f"thư mục tháng thiếu file: {', '.join(miss)}"]))
        if bad:
            out[sid] = bad
    os.makedirs(os.path.dirname(CACHE), exist_ok=True)
    io.open(CACHE, "w", encoding="utf-8").write(json.dumps(cache, ensure_ascii=False))
    return out


def main(argv):
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    res = validate_all(set(argv) or None)
    if not res:
        print("✔ Mọi file đúng mẫu.")
        return 0
    for sid, items in res.items():
        print(f"✖ {sid} — {SOURCES[sid]['name']}")
        for rel, errs in items:
            print(f"   {rel}")
            for e in errs:
                print(f"      · {e}")
    return 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
