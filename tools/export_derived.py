# -*- coding: utf-8 -*-
"""
XUẤT BẢNG LUỸ KẾ TỪ HAI LANE CŨ VÀO data_input/
===============================================
    python tools/export_derived.py          (update.py tự gọi sau build_hub/build_mkt)

Vì sao cần file này
-------------------
Dashboard đọc data_input/*.xlsx. `tools/build_month.py` dựng các bảng THEO THÁNG,
nhưng ~25 bảng luỹ kế (Menu · khung giờ · khu vực · nhân viên · CTKM · voucher · đối
tác · KPI CRM · _stats) chỉ có build_hub.py / build_mkt.py tính được — và hai lane đó
ghi ra data.json / data_mkt.json ở thư mục gốc mà KHÔNG AI ĐỌC. Kết quả: các bảng
này trong data_input/ đóng băng từ đợt di trú 10/09/2026, thả file mới bao nhiêu
cũng không đổi.

Nguyên tắc chống cộng đôi
-------------------------
Chỉ ghi những bảng KHÔNG có lane nào khác dựng (danh sách EXPORT dưới đây). Bảng mà
tools/build_month.py đã dựng (nature · ads_* · daily…) tuyệt đối không ghi ở đây.

Ghi TẠI CHỖ: chỉ thay đúng sheet được liệt kê, mọi sheet khác của file giữ nguyên —
kể cả dim_store (bí danh cửa hàng) và các sheet khai tay trong 01_master.xlsx.
Bảng nguồn rỗng (lane lỗi, thiếu file) thì BỎ QUA, không ghi đè bằng bảng rỗng.
"""
from __future__ import annotations

import io
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from monthly_lib import DATA_INPUT, MONTHLY_DIR, ROOT, cols_of  # noqa: E402

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

MASTER = os.path.join(DATA_INPUT, "01_master.xlsx")
SNAPSHOT = os.path.join(DATA_INPUT, "02_snapshot.xlsx")


def load(name):
    p = os.path.join(ROOT, name)
    if not os.path.exists(p):
        return {}
    return json.load(io.open(p, encoding="utf-8"))


def _cell(v):
    if isinstance(v, (list, dict)):
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, bool):
        return int(v)
    return v


def replace_sheets(path, tables):
    """Thay đúng các sheet trong `tables` = {tên: [dict]}; sheet khác giữ nguyên."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill

    tables = {k: v for k, v in tables.items() if v}
    if not tables:
        return 0
    wb = load_workbook(path) if os.path.exists(path) else Workbook()
    if not os.path.exists(path) and wb.active and wb.active.max_row <= 1:
        wb.remove(wb.active)
    for name, rows in tables.items():
        idx = wb.sheetnames.index(name) if name in wb.sheetnames else len(wb.sheetnames)
        if name in wb.sheetnames:
            wb.remove(wb[name])
        ws = wb.create_sheet(name, idx)
        want = cols_of(name) or []
        extra = [k for r in rows for k in r if k not in want]
        header = want + list(dict.fromkeys(extra))
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True, color="F3F2EE")
            c.fill = PatternFill("solid", fgColor="1F1F26")
        for r in rows:
            ws.append([_cell(r.get(h)) for h in header])
        ws.freeze_panes = "A2"
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return len(tables)


def by_month(rows, key):
    out = {}
    for r in rows or []:
        m = str(r.get(key) or "")[:7]
        if len(m) == 7:
            out.setdefault(m, []).append(r)
    return out


def stats_rows(H, M):
    """Dòng `_stats` mới — CHỈ các khối tính lại được từ lần chạy này."""
    blocks = {
        "product_stat": H.get("product_stat"),
        "repeat_stat": H.get("repeat_stat"),
        "menu_median": H.get("menu_median"),
        "gads_kw_stat": M.get("gads_kw_stat"),
        "voucher_stat": M.get("voucher_stat"),
        "member_stat": M.get("member_stat"),
    }
    vj = M.get("voucher_join")
    if isinstance(vj, dict):
        blocks["voucher_join"] = {k: v for k, v in vj.items() if k != "by_month"}
    bud = M.get("budget")
    if isinstance(bud, dict):
        blocks["budget"] = {k: bud[k] for k in ("total", "plan", "gap", "file") if k in bud}
    meta = H.get("meta") or {}
    if "cogs_coverage" in meta:
        blocks["meta"] = {"cogs_coverage": meta["cogs_coverage"]}
    ms = meta.get("months") or []
    if ms and isinstance(blocks.get("product_stat"), dict):
        blocks["product_stat"] = dict(blocks["product_stat"], covers=f"{ms[0]} → {ms[-1]}")
    rows = []
    for t, d in blocks.items():
        if isinstance(d, dict) and d:
            rows += [{"table": t, "field": f, "value": v} for f, v in d.items()]
    return rows


def merge_stats(path, fresh):
    """Giữ dòng _stats cũ của khối KHÔNG tính lại được, thay khối đã tính lại.
    `ads_stat` bị GỠ hẳn: bảng ads nay do build_month dựng theo tháng và loader tự
    cộng — giữ số khai cũ (197,9tr đóng băng) là đè lên số thật."""
    from openpyxl import load_workbook
    old = []
    if os.path.exists(path):
        wb = load_workbook(path, read_only=True, data_only=True)
        if "_stats" in wb.sheetnames:
            it = wb["_stats"].iter_rows(values_only=True)
            head = next(it, None)
            for r in it:
                if r and r[0]:
                    old.append({"table": r[0], "field": r[1], "value": r[2]})
        wb.close()
    redo = {r["table"] for r in fresh} | {"ads_stat"}
    keep = [r for r in old if r["table"] not in redo
            and not (r["table"] == "meta" and r["field"] == "source_root")]
    return keep + fresh


def main():
    H, M = load("data.json"), load("data_mkt.json")
    if not H and not M:
        print("  ✖ không có data.json / data_mkt.json — chạy build_hub.py, build_mkt.py trước")
        return 2
    print("─" * 74)
    print("XUẤT BẢNG LUỸ KẾ → data_input/")
    n = 0

    # ── 02_snapshot.xlsx ── bảng luỹ kế toàn kỳ
    dwell = []
    if isinstance(H.get("dwell"), dict) and H["dwell"].get("n"):
        dwell.append(dict(store=None, **{k: H["dwell"].get(k) for k in ("n", "mean", "median")}))
    dwell += [dict(store=r["store"], mean=r.get("dwell")) for r in H.get("dwell_store") or []]
    snap = {
        "product": H.get("product"), "category": H.get("category"), "group": H.get("group"),
        "heat": H.get("heat"), "zone": H.get("zone"), "staff": H.get("staff"),
        "payment": H.get("payment"), "repeat": H.get("repeat"), "campaigns": H.get("campaigns"),
        "dwell": dwell, "voucher_prog": M.get("voucher_prog"),
    }
    stats = stats_rows(H, M)
    if stats:
        snap["_stats"] = merge_stats(SNAPSHOT, stats)
    k = replace_sheets(SNAPSHOT, snap)
    print(f"  02_snapshot.xlsx   thay {k} sheet: {', '.join(x for x, v in snap.items() if v)}")
    n += k

    # ── 01_master.xlsx ── chỉ ba bảng máy tính được; ngân sách · dim_store · pre_analytics GIỮ tay
    mast = {"crm_target": M.get("crm_target"), "partners": M.get("partners"),
            "partner_camp": M.get("partner_camp")}
    k = replace_sheets(MASTER, mast)
    print(f"  01_master.xlsx     thay {k} sheet: {', '.join(x for x, v in mast.items() if v) or '—'}")
    n += k

    # ── monthly/YYYY-MM.xlsx ── voucher theo tháng + độ phủ giá vốn
    vm = by_month(M.get("voucher_month"), "m")
    vj = by_month((M.get("voucher_join") or {}).get("by_month") if isinstance(M.get("voucher_join"), dict) else [], "m")
    cc = by_month(H.get("cogs_cov"), "month")
    for month in sorted(set(vm) | set(vj) | set(cc)):
        t = {"voucher_month": vm.get(month),
             "voucher_join": [{k2: r.get(k2) for k2 in ("m", "n", "hit")} for r in vj.get(month, [])],
             "cogs_cov": cc.get(month)}
        k = replace_sheets(os.path.join(MONTHLY_DIR, f"{month}.xlsx"), t)
        n += k
    print(f"  monthly/*.xlsx     voucher_month {len(vm)} tháng · voucher_join {len(vj)} · cogs_cov {len(cc)}")
    print(f"  tổng {n} sheet được làm mới")
    return 0


if __name__ == "__main__":
    sys.exit(main())
