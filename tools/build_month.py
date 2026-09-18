# -*- coding: utf-8 -*-
"""
NOIRE ANALYTICS HUB — DỰNG GÓI SỐ LIỆU MỘT THÁNG TỪ DỮ LIỆU THÔ
================================================================
    python tools/build_month.py 2026-08          dựng một tháng
    python tools/build_month.py 2026-07 2026-08  dựng nhiều tháng
    python tools/build_month.py --all            dựng mọi tháng tìm thấy trong nguồn
    python tools/build_month.py 2026-08 --dry    chỉ in ra, không ghi file

Đầu ra: data_input/monthly/YYYY-MM.xlsx
Sheet nào script dựng được thì GHI ĐÈ; sheet còn lại trong file cũ được GIỮ NGUYÊN,
nên chạy lại không xoá mất phần nhập tay.

NGUỒN ĐỌC: L0_input/ theo sổ đăng ký data_sources.json (xem tools/l0_registry.py).
  python tools/build_month.py 2026-08 --parts=meta,oa   chỉ dựng lại vài phần
  Phần hợp lệ: tracking · pos · meta · google · social · oa · member · promotion ·
               booking · partnership
Thường không cần gọi trực tiếp — update.py tự chọn tháng + phần cần dựng.
"""
from __future__ import annotations

import glob
import os
import re
import sys
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from monthly_lib import (  # noqa: E402
    CONTRACT,
    BRAND_OF_STORE, ads_is_booking, ads_page, ads_result_kind, booking_lost_reason,
    booking_etype, booking_segment, booking_stage, L0_ROOT, L0_WHY, MONTHLY_DIR, ROOT, SOURCES, agg,
    classify_nature, date_of, days_in_month, l0_by_month, l0_dir, l0_latest,
    l0_files, l0_month_file, l0_month_files, month_of, norm,
    read_html_table, read_utf16_csv, read_workbook, store_code, store_in_text,
    to_num, write_workbook,
)
from split_to_monthly import guide_monthly  # noqa: E402

# ══════════════════════════════════════════════════════════════════════
# Gốc dữ liệu thô
# ══════════════════════════════════════════════════════════════════════
# Gốc dữ liệu = L0_input/ (monthly_lib.L0_ROOT). Trước 16/09/2026 mỗi script có
# find_root() riêng và chốt trên hai cây khác nhau.
RAW_ROOT = L0_ROOT if os.path.isdir(L0_ROOT) else None
RAW = RAW_ROOT
# Tracking Sales KHÔNG phải file thô: tools/tracking.py sinh nó từ file doanh thu
# ngày (S03_daily) + config_targets.csv (S00_targets) ở L0_input. Nằm trong _cache
# vì là sản phẩm trung gian — đặt trong data_input/ thì loader đọc nhầm thành sheet lạ.
TRACKING = os.path.join(ROOT, "_cache", "tracking", "NOIRE_Tracking_Sales_2026.xlsx")

WARN: list[str] = []


def log(*a):
    print(*a)


def warn(msg):
    WARN.append(msg)
    print(f"      ! {msg}")


def vn_month_dir(base, month):
    """Thư mục kiểu 'Tháng 8.2026'. Chấp nhận cả biến thể có/không số 0 đứng đầu."""
    y, m = month[:4], int(month[5:7])
    for pat in (f"Tháng {m}.{y}", f"Tháng {m:02d}.{y}", f"T{m}.{y}", f"{month}"):
        p = os.path.join(base, pat)
        if os.path.isdir(p):
            return p
    return None


def sheet_rows(path, sheet=None, header_row=1, reset=False):
    """Đọc một sheet thành list[dict]; header_row là dòng chứa tên cột (1-based)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    if reset:
        ws.reset_dimensions()
    rows, header = [], None
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < header_row:
            continue
        if i == header_row:
            header = [str(c).strip() if c is not None else f"_c{j}" for j, c in enumerate(r)]
            continue
        if r is None or not any(c is not None and str(c).strip() for c in r):
            continue
        rows.append({header[j]: r[j] for j in range(min(len(header), len(r)))})
    wb.close()
    return rows


def col(d: dict, *needles, exact=None):
    """Tìm tên cột theo từ khoá đã chuẩn hoá — file export hay đổi hoa/thường và dấu."""
    if exact and exact in d:
        return exact
    keys = list(d.keys())
    for n in needles:
        nn = norm(n)
        for k in keys:
            if nn in norm(k):
                return k
    return None


# ══════════════════════════════════════════════════════════════════════
# 1. TRACKING SALES TOOL → store_month · daily · coverage · dim_target
# ══════════════════════════════════════════════════════════════════════
def read_tracking(month):
    if not TRACKING or not os.path.exists(TRACKING):
        warn("chưa có file Tracking Sales — chạy python tools/tracking.py (update.py tự chạy)")
        return {}

    rows = sheet_rows(TRACKING, "Data_Daily", header_row=4)
    daily, unknown = [], set()
    for r in rows:
        d = date_of(r.get("Ngày"))
        if not d or d[:7] != month:
            continue
        code = store_code(r.get("Cửa hàng (export)"))
        if not code:
            unknown.add(str(r.get("Cửa hàng (export)")))
            continue
        daily.append({
            "date": d, "store": code,
            "net": to_num(r.get("Net Sales"), 0),
            "gross": to_num(r.get("Gross Sales"), 0),
            "guest": to_num(r.get("GC (Guest)"), 0),
            "tc": to_num(r.get("TC"), 0),
        })
    if unknown:
        warn("cửa hàng chưa khai ở dim_store.aliases (01_master.xlsx): %s"
             % ", ".join(sorted(unknown)))
    if not daily:
        warn(f"Tracking Sales không có dòng nào của {month}")
        return {}

    sm = agg(daily, ["store"], sums=["net", "gross", "guest", "tc"])
    store_month = []
    for r in sorted(sm, key=lambda x: x["store"]):
        net, gross = r["net"], r["gross"]
        # Cửa hàng chưa khai trương vẫn có dòng toàn số 0 trong file tracking.
        # Nộp lên sẽ vẽ ra một tháng kinh doanh không có thật và làm hỏng chốt #5.
        if not net and not r["tc"] and not r["guest"]:
            continue
        store_month.append({
            "month": month, "store": r["store"],
            "net": round(net), "guest": round(r["guest"]), "tc": round(r["tc"]),
            "gross": round(gross),
            # disc = gross − net. Đây là ĐỊNH NGHĨA đang dùng của hệ thống, đã đối
            # soát khớp bản do lane POS sinh ra cho T1–T7.
            "disc": round(gross - net) if gross else None,
        })

    dates = sorted({r["date"] for r in daily})
    coverage = [{
        "month": month, "days_data": len(dates), "days_month": days_in_month(month),
        "first": dates[0], "last": dates[-1],
    }]

    out = {
        "store_month": store_month,
        # Ngày cửa hàng đóng / chưa mở: bỏ, đỡ phình file mà không mất thông tin —
        # số ngày có dữ liệu của tháng nằm ở sheet coverage.
        "daily": [{k: v for k, v in r.items() if k in ("date", "store", "net", "guest", "tc")}
                  for r in sorted(daily, key=lambda x: (x["date"], x["store"]))
                  if r["net"] or r["tc"] or r["guest"]],
        "coverage": coverage,
    }
    tgt = read_tracking_target(month)
    if tgt:
        out["dim_target"] = tgt
    return out


def read_tracking_target(month):
    """Bóc dòng `Target` của sheet `By Month`. Chỉ lấy các block ánh xạ được về
    một cửa hàng — block `Others · 9 Stellars + Gateway` gộp ba cửa hàng nên bỏ,
    target của chúng đã khai sẵn ở 01_master.xlsx."""
    try:
        rows = sheet_rows(TRACKING, "By Month", header_row=4)
    except KeyError:
        return []
    ti = f"T{int(month[5:7])}"
    out, cur = [], None
    for r in rows:
        vals = list(r.values())
        grp, store, metric = (vals + [None, None, None])[:3]
        if store and str(store).strip():
            nm = str(store).strip()
            # Bỏ dòng brand ("NCB · Noire Café & Bistro") và dòng gộp nhiều cửa hàng.
            cur = None if (" · " in nm and not nm.startswith("IFC")) or "+" in nm \
                else store_code(nm)
        if cur and metric and norm(metric) == "target":
            v = to_num(r.get(ti))
            if v is not None:
                out.append({"month": month, "store": cur, "target": round(v)})
    return out


# ══════════════════════════════════════════════════════════════════════
# 2. META ADS → ads_month · ads_brand · ads_objective · ads_campaign_detail
# ══════════════════════════════════════════════════════════════════════
BRAND_PAT = [("NJFB", r"\bnjfb\b|\bnfb\b|\bjfb\b|japanese|fusion|crest|\bssv\b"),
             ("NDC", r"\bndc\b|dining|39 ?ntmk|berkley"),
             ("NCB", r"\bncb\b|bistro|the mett|empress|\bskc\b|café|cafe")]
# Chiến dịch tuyển dụng KHÔNG phải marketing thương hiệu — phải tách khỏi Ad Cost Ratio.
HR_PAT = r"hr |tuyển dụng|tuyen dung|recruit"
OBJ_PAT = [("Tuyển dụng", HR_PAT),
           ("Lead tiệc", r"\blead\b|tiệc|tiec|\byep\b|booking|party|banquet"),
           ("Tin nhắn", r"message|mess\b|tin nhắn|cuộc trò chuyện|inbox"),
           ("Tương tác", r"engagement|tương tác|post|reel|clip|video"),
           ("Tiếp cận", r"reach|awareness|nhận diện|traffic|lượt xem")]


def tag(txt, pats, default=None):
    n = norm(txt)
    for lab, p in pats:
        if re.search(p, n):
            return lab
    return default


def read_meta_ads(month):
    hit = l0_month_file("S08_ads_meta", month)
    files = [hit] if hit else []
    if not files:
        warn(f"không thấy báo cáo Meta Ads của {month} trong {SOURCES['S08_ads_meta']['dir']}")
        return {}

    camps = []
    for f in sorted(files):
        try:
            rows = sheet_rows(f, "Raw Data Report", header_row=3)
        except KeyError:
            rows = sheet_rows(f, header_row=3)
        if not rows:
            continue
        c0 = rows[0]
        k_camp = col(c0, "tên chiến dịch")
        k_lvl = col(c0, "cấp độ phân phối")
        k_res = col(c0, exact="Kết quả") or col(c0, "kết quả")
        k_rty = col(c0, "loại kết quả")
        k_spd = col(c0, "số tiền đã chi tiêu", "amount spent")
        k_imp = col(c0, "lượt hiển thị", "impression")
        k_rch = col(c0, "người tiếp cận", "reach")
        # BẪY: khớp chuỗi con "lượt click vào liên kết" trúng cột CPC đứng TRƯỚC
        # ("CPC (chi phí trên mỗi lượt click vào liên kết)") — ra đơn giá, không ra số click.
        k_clk = col(c0, "link clicks", exact="Lượt click vào liên kết")
        if not k_camp:
            warn(f"{os.path.basename(f)}: không thấy cột 'Tên chiến dịch'")
            continue

        n_all = 0
        for r in rows:
            if not r.get(k_camp):
                continue
            n_all += 1
            # BẪY: export có cả dòng cấp 'campaign' lẫn 'adset', vài tháng còn tách
            # theo tuổi × giới tính. Cộng tất cả ra gấp 3,5 lần số thật.
            if k_lvl and norm(r.get(k_lvl)) != "campaign":
                continue
            camp = str(r[k_camp]).strip()
            is_hr = bool(re.search(HR_PAT, norm(camp)))
            camps.append({
                "campaign": camp,
                "brand": "Tuyển dụng" if is_hr else tag(camp, BRAND_PAT, "Không xác định"),
                "objective": tag(f"{camp} {r.get(k_rty, '')}", OBJ_PAT, "Khác"),
                "result": to_num(r.get(k_res), 0) if k_res else 0,
                "spend": to_num(r.get(k_spd), 0) if k_spd else 0,
                "impr": to_num(r.get(k_imp), 0) if k_imp else 0,
                "reach": to_num(r.get(k_rch), 0) if k_rch else 0,
                "clicks": to_num(r.get(k_clk), 0) if k_clk else 0,
                # Phễu booking M10: nhãn gắn ở ĐÂY theo luật $booking của hợp đồng,
                # loader và màn hình chỉ đọc nhãn.
                "page": ads_page(camp),
                "rkind": ads_result_kind(r.get(k_rty)),
                "funnel": "booking" if not is_hr and ads_is_booking(camp) else None,
            })
        log(f"      Meta: {os.path.basename(f)} — {n_all} dòng → {len(camps)} cấp campaign")

    if not camps:
        return {}

    det = agg(camps, ["campaign", "brand", "objective", "page", "rkind", "funnel"],
              sums=["spend", "result", "reach", "impr", "clicks"])
    det = sorted(det, key=lambda r: -r["spend"])
    return {
        "ads_month": [{
            "month": month,
            "spend": round(sum(c["spend"] for c in camps)),
            "reach": round(sum(c["reach"] for c in camps)),
            "impr": round(sum(c["impr"] for c in camps)),
            "n": len({c["campaign"] for c in camps}),
        }],
        "ads_brand": [{"month": month, "brand": r["brand"],
                       "spend": round(r["spend"]), "reach": round(r["reach"])}
                      for r in sorted(agg(camps, ["brand"], sums=["spend", "reach"]),
                                      key=lambda x: -x["spend"])],
        "ads_objective": [{"month": month, "objective": r["objective"],
                           "spend": round(r["spend"]), "result": round(r["result"])}
                          for r in sorted(agg(camps, ["objective"], sums=["spend", "result"]),
                                          key=lambda x: -x["spend"])],
        "ads_campaign_detail": [{
            "month": month, "brand": r["brand"], "objective": r["objective"],
            "campaign": r["campaign"], "spend": round(r["spend"]),
            "reach": round(r["reach"]), "impr": round(r["impr"]), "result": round(r["result"]),
            "page": r["page"], "rkind": r["rkind"], "funnel": r["funnel"],
            "clicks": round(r["clicks"]),
        } for r in det],
    }


# ══════════════════════════════════════════════════════════════════════
# 3. GOOGLE ADS → ads_google · gads_channel · gads_kw
# ══════════════════════════════════════════════════════════════════════
# BẪY: mọi báo cáo Google đều xen dòng cộng dồn "Tổng số: …". Không loại thì
# đứng đầu mọi bảng xếp hạng là mấy dòng tổng, và chi phí bị nhân đôi.
IS_TOTAL = re.compile(r"^\s*tổng số\s*:", re.I)

# Bảng nhận cửa hàng từ tên chiến dịch đã chuyển về cột `alias_re` của sheet
# dim_store (01_master.xlsx), dùng qua monthly_lib.store_in_text(). Trước đây file
# này và build_mkt.py mỗi bên giữ một bản — và hai bản ĐÃ lệch: bản kia nhận thêm
# `minh khai` cho NDC_NTMK, bản này thì không.


def read_google_ads(month):
    base = l0_dir("S09_ads_google")
    if not base:
        return {}
    folder = vn_month_dir(base, month)
    if not folder:
        warn(f"không thấy thư mục Google Ads của {month}")
        return {}
    out = {}

    # ---- chiến dịch ----
    p = os.path.join(folder, "Báo cáo chiến dịch.xlsx")
    if os.path.exists(p):
        rows = sheet_rows(p, header_row=3)
        c0 = rows[0] if rows else {}
        k_st = col(c0, "trạng thái chiến dịch")
        k_nm = col(c0, exact="Chiến dịch") or col(c0, "chiến dịch")
        k_bd = col(c0, "ngân sách")
        k_sp = col(c0, exact="Chi phí") or col(c0, "chi phí")
        k_cv = col(c0, "lượt chuyển đổi")
        k_ck = col(c0, "lượt nhấp")
        k_im = col(c0, "số lượt hiển thị")
        g = []
        for r in rows:
            nm = r.get(k_nm)
            if not nm or IS_TOTAL.match(str(r.get(k_st) or "")) or IS_TOTAL.match(str(nm)):
                continue
            nn = norm(nm)
            st = store_in_text(nn)
            g.append({
                "month": month, "campaign": str(nm).strip(), "store": st,
                "brand": tag(nm, BRAND_PAT, "Không xác định"),
                "status": str(r.get(k_st) or "").strip(),
                "budget_day": to_num(r.get(k_bd)), "spend": to_num(r.get(k_sp), 0),
                "conv": to_num(r.get(k_cv), 0), "clicks": to_num(r.get(k_ck), 0),
                "impr": to_num(r.get(k_im), 0),
            })
        if g:
            out["ads_google"] = sorted(g, key=lambda x: -x["spend"])
            log(f"      Google: {len(g)} chiến dịch · chi {sum(x['spend'] for x in g):,.0f} đ"
                .replace(",", "."))

    # ---- kênh phân phối ----
    p = os.path.join(folder, "Báo cáo thông tin chi tiết về cụm từ tìm kiếm.xlsx")
    if os.path.exists(p):
        rows = sheet_rows(p, header_row=3)
        c0 = rows[0] if rows else {}
        k_ch = col(c0, exact="Kênh") or col(c0, "kênh")
        k_im = col(c0, "số lượt hiển thị")
        k_ck = col(c0, "lượt nhấp")
        k_cv = col(c0, "lượt chuyển đổi")
        k_sp = col(c0, exact="Chi phí") or col(c0, "chi phí")
        ch = [{"channel": str(r[k_ch]).strip(), "impr": to_num(r.get(k_im), 0),
               "clicks": to_num(r.get(k_ck), 0), "conv": to_num(r.get(k_cv), 0),
               "spend": to_num(r.get(k_sp), 0)}
              for r in rows if r.get(k_ch) and not IS_TOTAL.match(str(r[k_ch]))]
        if ch:
            out["gads_channel"] = [
                {"month": month, **r} for r in
                sorted(agg(ch, ["channel"], sums=["impr", "clicks", "conv", "spend"]),
                       key=lambda x: -x["spend"])]

    # ---- cụm từ tìm kiếm ----
    p = os.path.join(folder, "Báo cáo cụm từ tìm kiếm.xlsx")
    if os.path.exists(p):
        rows = sheet_rows(p, header_row=3)
        c0 = rows[0] if rows else {}
        k_kw = col(c0, "cụm từ tìm kiếm")
        k_ck = col(c0, "lượt nhấp")
        k_im = col(c0, "số lượt hiển thị")
        k_sp = col(c0, exact="Chi phí") or col(c0, "chi phí")
        k_cv = col(c0, "lượt chuyển đổi")
        kw = [{"kw": str(r[k_kw]).strip(), "clicks": to_num(r.get(k_ck), 0),
               "impr": to_num(r.get(k_im), 0), "spend": to_num(r.get(k_sp), 0),
               "conv": to_num(r.get(k_cv), 0)}
              for r in rows if r.get(k_kw) and not IS_TOTAL.match(str(r[k_kw]))]
        kw = agg(kw, ["kw"], sums=["clicks", "impr", "spend", "conv"])
        # Giữ cụm từ CÓ HOẠT ĐỘNG. Phần đuôi chỉ có impression bằng 0 chiếm ~90% số
        # dòng mà không nói lên điều gì — bỏ để file tháng không phình lên vài nghìn dòng.
        act = [r for r in kw if r["clicks"] or r["conv"] or r["spend"]]
        if act:
            out["gads_kw"] = [{"month": month, **r}
                              for r in sorted(act, key=lambda x: -x["clicks"])]
            log(f"      Google: {len(act)}/{len(kw)} cụm từ có hoạt động")
    return out


# ══════════════════════════════════════════════════════════════════════
# 4. FANPAGE FACEBOOK → social_month
# ══════════════════════════════════════════════════════════════════════
# Thư mục brand trong export ↔ (brand hệ thống, tên kênh hiển thị)
FB_PAGES = {
    "NCB": ("NCB", "NOIRE Café & Bistro"),
    "NDC": ("NDC", "NOIRE Dining & Cafe"),
    "NJFB": ("NJFB", "NOIRE Japanese Fusion & Bar"),
    # NOIRE Express · Creative Park là fanpage riêng, không thuộc ba brand chính.
    "NEC": ("OTHER", "NOIRE Express · Creative Park"),
}
# Tên file CSV ↔ trường trong social_month. Facebook đặt tên theo tiếng Việt.
FB_METRIC = {
    "người xem": "reach",            # số TÀI KHOẢN tiếp cận
    "lượt xem": "views",             # số LƯỢT xem — khác reach, không cộng chung
    "lượt tương tác": "engage",
    "lượt theo dõi": "follows",
    "lượt truy cập": "profile_views",
    "lượt click vào liên kết": "clicks",
}


FB_COLS = {
    "month": ("tháng", "month"),
    "page": ("fanpage", "trang", "page"),
    "views": ("lượt xem", "views"),
    "reach": ("người xem trong kỳ", "người xem", "reach"),
    "engage": ("lượt tương tác với nội dung", "lượt tương tác", "engagement"),
    "clicks": ("lượt click vào liên kết", "link clicks"),
    "contacts": ("tổng số người liên hệ", "người liên hệ", "contacts"),
    "msgs": ("lượt bắt đầu cuộc trò chuyện qua tin nhắn", "messaging conversations started"),
    "profile_views": ("lượt truy cập", "page visits"),
    "follows": ("lượt theo dõi", "follows"),
    "unfollows": ("lượt bỏ theo dõi", "unfollows"),
    "posts": ("số bài đăng", "posts"),
    "spend": ("chi phí boost", "spend"),
}


def _fb_from_summary(month):
    """File `Facebook*.xlsx` (sheet Tong_hop_thang). Nhiều file thì file sửa SAU thắng
    cho cùng tháng × fanpage — thả bản T7–T9 mới không cần xoá bản T7–T8 cũ."""
    files = [f for f in l0_files("S18_social") if f.lower().endswith((".xlsx", ".xlsm"))]
    got = {}
    for path in sorted(files, key=os.path.getmtime):
        sheet, rows = find_table(path, FB_COLS, ["month", "page", "views"])
        if not sheet:
            warn(f"{os.path.basename(path)}: không thấy bảng Tháng · Fanpage · Lượt xem")
            continue
        for r in rows:
            if month_of(r.get("month")) != month:
                continue
            code = str(r.get("page") or "").strip().upper()
            brand, page = FB_PAGES.get(code, ("OTHER", str(r.get("page") or "").strip()))
            rec = {"month": month, "platform": "FACEBOOK", "brand": brand, "page": page,
                   "code": code or None}
            for k in ("views", "reach", "engage", "clicks", "contacts", "msgs", "profile_views",
                      "follows", "unfollows", "posts", "spend"):
                v = to_num(r.get(k))
                if v is not None:
                    rec[k] = round(v)
            rec["days"] = days_in_month(month)
            got[(brand, page)] = rec
    return list(got.values())


def read_facebook_pages(month):
    rows = _fb_from_summary(month)
    if rows:
        log(f"      Fanpage (file tổng hợp): {len(rows)} kênh · "
            f"{sum(r.get('reach', 0) for r in rows):,} người xem trong kỳ".replace(",", "."))
        return {"social_month": rows}
    # ---- dự phòng: CSV theo ngày ----
    base = l0_dir("S18_social")
    folder = vn_month_dir(base, month) if base else None
    if not folder:
        return {}
    rows = []
    for d in sorted(os.listdir(folder)):
        sub = os.path.join(folder, d)
        if not os.path.isdir(sub):
            continue
        brand, page = FB_PAGES.get(d.upper(), ("OTHER", d))
        rec = {"month": month, "platform": "FACEBOOK", "brand": brand, "page": page,
               "code": d.upper()}
        days = 0
        for f in glob.glob(os.path.join(sub, "*.csv")):
            key = FB_METRIC.get(norm(os.path.splitext(os.path.basename(f))[0]))
            # BẪY: "Người xem" là số NGƯỜI duy nhất — cộng 31 ngày là đếm một người
            # nhiều lần (NCB T8: 322.042 cộng ngày so với 238.590 cả kỳ). Bỏ trống.
            if not key or key == "reach":
                continue
            try:
                _, header, data = read_utf16_csv(f)
            except Exception as e:                                  # noqa: BLE001
                warn(f"{d}/{os.path.basename(f)}: {str(e)[:60]}")
                continue
            vals = [to_num(r[1]) for r in data if len(r) > 1 and to_num(r[1]) is not None]
            if not vals:
                continue
            rec[key] = round(sum(vals))
            days = max(days, len(vals))
        if len(rec) > 5:                     # month · platform · brand · page · code + ≥1 chỉ số
            rec["days"] = days
            rows.append(rec)
    if rows:
        warn(f"Fanpage {month}: chỉ có CSV theo ngày — reach để TRỐNG (không cộng theo ngày được). "
             "Nhập file tổng hợp theo mẫu _MAU_Facebook_Tong_hop.xlsx để có reach.")
        return {"social_month": rows}
    return {}


# TikTok — nguồn DUY NHẤT hợp lệ là file tổng hợp theo mẫu `_MAU_TikTok_Tong_hop.xlsx`
# (TikTok Studio không xuất file số tháng, team chép số từ ảnh chụp) hoặc file Overview
# tải từ TikTok Studio. Nhận diện cột theo TỪ KHOÁ, cả tiếng Việt lẫn tiếng Anh.
TT_COLS = {
    "month": ("tháng", "month"),
    "date": ("ngày", "date"),
    "page": ("tài khoản", "account"),
    "views": ("lượt xem bài đăng", "lượt xem video", "video views", "post views"),
    "profile_views": ("lượt xem hồ sơ", "profile views"),
    "likes": ("thích", "likes"),
    "comments": ("bình luận", "comments"),
    "shares": ("chia sẻ", "shares"),
    "net_follow": ("follower ròng", "net followers", "follower thuần"),
    "followers": ("tổng follower", "total followers", "người theo dõi"),
    "posts": ("số bài", "videos posted", "bài đăng mới"),
}
TT_DEFAULT_PAGE = "NOIRE F&B (TikTok)"


def find_table(path, colmap, must, prefer_sheet="tong_hop_thang"):
    """Tìm bảng số trong một file tổng hợp: sheet nào có dòng tiêu đề chứa đủ các cột
    `must`. `colmap` = {trường: (từ khoá, …)} — khớp CHÍNH XÁC trước, rồi mới khớp chuỗi
    con (từ khoá > 5 ký tự), để `lượt xem` không nuốt `lượt xem hồ sơ`.
    Tiêu đề thường KHÔNG ở dòng 1 (có tựa lớn + dòng trống) nên dò theo nội dung."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    names = sorted(wb.sheetnames, key=lambda n: norm(n) != prefer_sheet)
    try:
        for nm in names:
            ws = wb[nm]
            ws.reset_dimensions()
            hdr, rows = None, []
            for r in ws.iter_rows(values_only=True):
                if not r:
                    continue
                if hdr is None:
                    cells = [norm(c) if c is not None else "" for c in r]
                    idx, used = {}, set()
                    for exact in (True, False):
                        for key, needles in colmap.items():
                            if key in idx:
                                continue
                            for j, c in enumerate(cells):
                                if not c or j in used:
                                    continue
                                if any(c == n if exact else (len(n) > 5 and n in c) for n in needles):
                                    idx[key] = j
                                    used.add(j)
                                    break
                    if all(k in idx for k in must):
                        hdr = idx
                    continue
                rec = {k: (r[j] if j < len(r) else None) for k, j in hdr.items()}
                if all(rec.get(k) is None for k in must):
                    continue
                rows.append(rec)
            if hdr:
                return nm, rows
    finally:
        wb.close()
    return None, []


def read_tiktok(month):
    path = l0_latest("S22_tiktok")
    if not path:
        warn("chưa có file TikTok (03_MARKETING/03_Social/TikTok) — dùng mẫu _MAU_TikTok_Tong_hop.xlsx")
        return []
    sheet, rows = find_table(path, TT_COLS, ["views"])
    rows = [r for r in rows if r.get("month") is not None or r.get("date") is not None]
    if not sheet:
        warn(f"{os.path.basename(path)}: không thấy bảng có cột 'Lượt xem' + 'Tháng'/'Ngày'")
        return []
    sel = [r for r in rows
           if (month_of(r.get("month")) or (date_of(r.get("date")) or "")[:7]) == month]
    if not sel:
        return []
    by_page = defaultdict(list)
    for r in sel:
        pg = str(r.get("page") or "").strip()
        if not pg or "chưa có tên" in norm(pg):
            pg = TT_DEFAULT_PAGE
        by_page[pg].append(r)
    out = []
    for pg, rs in by_page.items():
        def tot(k):
            v = [to_num(x.get(k)) for x in rs]
            v = [x for x in v if x is not None]
            return round(sum(v)) if v else None
        net = tot("net_follow")
        followers = [to_num(x.get("followers")) for x in rs if to_num(x.get("followers")) is not None]
        out.append({
            "month": month, "platform": "TIKTOK", "brand": "OTHER", "page": pg,
            "views": tot("views"), "profile_views": tot("profile_views"),
            "likes": tot("likes"), "comments": tot("comments"), "shares": tot("shares"),
            "posts": tot("posts"), "followers": round(followers[-1]) if followers else None,
            # TikTok chỉ cho số follower RÒNG. Ghi vào `follows` với `unfollows` = 0 để
            # net_follow = follows − unfollows vẫn đúng; KHÔNG đọc `follows` TikTok như
            # lượt theo dõi mới của Facebook.
            "follows": net, "unfollows": 0 if net is not None else None,
            "days": len(rs) if "date" in rs[0] and rs[0].get("date") is not None else None,
        })
    log(f"      TikTok: {len(out)} tài khoản · "
        f"{sum(r['views'] or 0 for r in out):,} lượt xem ({sheet})".replace(",", "."))
    return out


def read_social(month):
    """Social = Facebook (CSV Meta Business Suite) + TikTok. Một bảng social_month,
    `platform` luôn nằm trong khoá — loader không bao giờ cộng chéo hai nền tảng."""
    fb = (read_facebook_pages(month) or {}).get("social_month", [])
    tt = read_tiktok(month)
    rows = fb + tt
    return {"social_month": rows} if rows else {}


# ══════════════════════════════════════════════════════════════════════
# 5. ZALO OA → oa
# ══════════════════════════════════════════════════════════════════════
OA_METRIC = {
    "xem trang thông tin oa": "views",
    "gửi tin nhắn đến oa": "msgs",
    "tương tác thanh menu": "menu",
    "quan tâm": "follows",
    "xem nội dung": "content",
}


def read_oa(month):
    base = l0_dir("S12_zalo_oa")
    if not base:
        return {}
    hit = l0_month_file("S12_zalo_oa", month)
    if not hit:
        warn(f"không thấy file OA Zalo của {month}")
        return {}
    # BẪY: file mang đuôi .xls nhưng ruột là HTML — mở bằng openpyxl sẽ ném lỗi.
    header, data = read_html_table(hit)
    idx = {}
    for j, h in enumerate(header):
        k = OA_METRIC.get(norm(h))
        if k:
            idx[k] = j
    if not idx:
        warn(f"{os.path.basename(hit)}: không nhận ra cột nào")
        return {}
    rec = {"month": month, "days": len(data)}
    for k, j in idx.items():
        rec[k] = round(sum(to_num(r[j], 0) or 0 for r in data if len(r) > j))
    log(f"      Zalo OA: {rec['days']} ngày · quan tâm {rec.get('follows', 0)}")
    return {"oa": [rec]}


# ══════════════════════════════════════════════════════════════════════
# 6. MEMBER ĐĂNG KÝ → member
# ══════════════════════════════════════════════════════════════════════
def _crm_member(path, month):
    """CRM_Dashboard*.xlsx — sheet `KPI_Thang` (khách đăng ký/tháng) + sheet
    `Nguon_DangKy_*` (đăng ký + OA follow theo ngày). Tiêu đề không nằm ở dòng 1
    (có tiêu đề lớn + dòng trống phía trên) nên dò theo nội dung."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    member = days = oa = None
    try:
        if "KPI_Thang" in wb.sheetnames:
            ws = wb["KPI_Thang"]
            ws.reset_dimensions()
            hdr = None
            for r in ws.iter_rows(values_only=True):
                if not r:
                    continue
                cells = [norm(c) if c is not None else "" for c in r]
                if hdr is None:
                    if "tháng" in cells and "khách đăng ký" in cells:
                        hdr = (cells.index("tháng"), cells.index("khách đăng ký"))
                    continue
                if month_of(r[hdr[0]]) == month:
                    member = to_num(r[hdr[1]])
                    break
        # theo ngày: số ngày có số + OA follow
        for nm in [n for n in wb.sheetnames if norm(n).startswith("nguon_dangky")] + \
                  [n for n in wb.sheetnames if norm(n) == "dangky_ngay"]:
            ws = wb[nm]
            ws.reset_dimensions()
            hdr = None
            dd, oav = set(), []
            for r in ws.iter_rows(values_only=True):
                if not r:
                    continue
                cells = [norm(c) if c is not None else "" for c in r]
                # Một sheet có THỂ chứa nhiều bảng xếp chồng (tóm tắt ở trên, chi tiết
                # theo ngày ở dưới) — gặp dòng tiêu đề mới thì chuyển sang bảng đó.
                if "ngày" in cells and any("đăng ký" in c for c in cells):
                    jd = cells.index("ngày")
                    jo = next((j for j, c in enumerate(cells) if "oa follow" in c), None)
                    hdr = (jd, jo)
                    continue
                if hdr is None:
                    continue
                d = date_of(r[hdr[0]])
                if d and d[:7] == month:
                    dd.add(d)
                    if hdr[1] is not None and to_num(r[hdr[1]]) is not None:
                        oav.append(to_num(r[hdr[1]]))
            if dd:
                days = len(dd)
                # OA toàn cột 0/rỗng = CHƯA NHẬP, không phải bằng 0.
                oa = round(sum(oav)) if any(oav) else None
                break
    finally:
        wb.close()
    return member, oa, days


def read_member(month):
    files = l0_files("S13_member")
    crm = [f for f in files if os.path.basename(f).upper().startswith("CRM_DASHBOARD")]
    act = [f for f in files if os.path.basename(f).lower().startswith("member_actual")]
    if crm:
        path = max(crm, key=os.path.getmtime)
        member, oa, days = _crm_member(path, month)
        # Tháng không có dòng ở KPI_Thang = file CHƯA phủ tháng đó → không ghi gì.
        # (Bản thử đầu ghi `member=None, days=30` cho T9 vì bảng phụ có cột ngày lạ.)
        if member is None:
            return {}
        rec = {"month": month, "member": None if member is None else round(member),
               "oa": oa, "days": days}
        log(f"      Member (CRM Dashboard): {rec['member']} member mới · "
            f"{'—' if oa is None else oa} OA follow · {days or '—'} ngày chi tiết")
        return {"member": [rec]}
    if not act:
        warn("không thấy file CRM_Dashboard*.xlsx / member_actual*.xlsx")
        return {}
    path = max(act, key=os.path.getmtime)
    rows = sheet_rows(path, "Actual", header_row=1)
    c0 = rows[0] if rows else {}
    k_d = col(c0, exact="Ngày") or col(c0, "ngày")
    k_m = col(c0, "member đăng ký")
    k_o = col(c0, "oa follow")
    sel = [r for r in rows if (date_of(r.get(k_d)) or "")[:7] == month]
    if not sel:
        return {}

    def tot(k):
        """Ô trống ≠ 0: cả tháng không ai điền thì trả None, không trả 0."""
        vals = [to_num(r.get(k)) for r in sel]
        vals = [v for v in vals if v is not None]
        return round(sum(vals)) if vals else None

    rec = {"month": month, "member": tot(k_m), "oa": tot(k_o), "days": len(sel)}
    log(f"      Member: {rec['member']} member mới · "
        f"{'—' if rec['oa'] is None else rec['oa']} OA follow · {rec['days']} ngày")
    return {"member": [rec]}


# ══════════════════════════════════════════════════════════════════════
# 7. BÁO CÁO PROMOTION-AGG → aggregator · budget_nonmedia
# ══════════════════════════════════════════════════════════════════════
def _find_section(rows, marker, header_needle):
    """Báo cáo là biểu mẫu trình bày, không phải bảng phẳng: định vị theo CHỮ
    của mục ('C2.', 'C4.') rồi lấy dòng tiêu đề ngay dưới. Bám theo số dòng cứng
    sẽ vỡ ngay lần đầu ai đó chèn thêm một dòng."""
    start = None
    for i, r in enumerate(rows):
        txt = " ".join(str(c) for c in r if c is not None)
        if start is None and norm(txt).startswith(norm(marker)):
            start = i
        elif start is not None and norm(header_needle) in norm(txt):
            return i
    return None


def _grid(path, sheet):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    g = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return g


AGG_PLATFORM = {
    "grabfood": "GrabFood", "shopeefood": "ShopeeFood", "befood": "beFood",
    "dingning city": "Dining City", "dining city": "Dining City",
    "grab dine out": "Grab Dine Out",
}


def read_promotion(month):
    hit = l0_month_file("S19_aggregator", month)
    if not hit:
        warn(f"không thấy báo cáo Promotion-AGG của {month}")
        return {}
    out = {}

    # ---- C2 · theo nền tảng ----
    try:
        g = _grid(hit, "Aggregator")
    except KeyError:
        g = None
    if g:
        h = _find_section(g, "C2.", "Nền tảng")
        rows = []
        if h is not None:
            head = [norm(c) for c in g[h]]
            def at(r, *needles):
                for n in needles:
                    for j, c in enumerate(head):
                        if c and norm(n) in c:
                            return r[j] if j < len(r) else None
                return None
            cur_platform = None
            for r in g[h + 1:]:
                name = str(r[0]).strip() if r and r[0] else ""
                if not name:
                    continue
                if norm(name) in ("tổng", "tong", "khác", "khac"):
                    if norm(name).startswith("t"):
                        break
                    continue
                plat = AGG_PLATFORM.get(norm(name))
                st = store_code(name)
                if plat:
                    cur_platform = plat
                    # Dòng nền tảng có thể là TỔNG của các dòng cửa hàng ngay dưới.
                    # Giữ lại tạm, cuối vòng mới quyết định lấy dòng nào.
                    rows.append({"_lvl": "platform", "platform": plat, "store": None,
                                 "sales": to_num(at(r, "sales")),
                                 "orders": to_num(at(r, "order")),
                                 "items": to_num(at(r, "số món")),
                                 "guests": to_num(at(r, "khách")),
                                 "discount": to_num(at(r, "discount")),
                                 "commission": to_num(at(r, "commission")),
                                 "ads_spend": to_num(at(r, "ads spend")),
                                 "note": None})
                elif st and cur_platform:
                    rows.append({"_lvl": "store", "platform": cur_platform, "store": st,
                                 "sales": to_num(at(r, "sales")),
                                 "orders": to_num(at(r, "order")),
                                 "items": to_num(at(r, "số món")),
                                 "guests": to_num(at(r, "khách")),
                                 "discount": to_num(at(r, "discount")),
                                 "commission": to_num(at(r, "commission")),
                                 "ads_spend": to_num(at(r, "ads spend")),
                                 "note": None})
        # Nền tảng nào đã có dòng cửa hàng thì BỎ dòng tổng của nền tảng đó,
        # nếu không doanh thu bị đếm hai lần.
        has_store = {r["platform"] for r in rows if r["_lvl"] == "store"}
        keep = [r for r in rows if not (r["_lvl"] == "platform" and r["platform"] in has_store)]
        agg_rows = []
        for r in keep:
            if not r.get("sales"):
                continue
            code = r["store"]
            agg_rows.append({
                "month": month, "platform": r["platform"],
                "brand": (code or "")[:4].rstrip("_") or None,
                "store": code,
                "sales": round(r["sales"]), "orders": round(r["orders"] or 0),
                "items": r["items"], "guests": r["guests"],
                "discount": r["discount"], "commission": r["commission"],
                "ads_spend": r["ads_spend"], "note": r["note"],
            })
        if agg_rows:
            out["aggregator"] = agg_rows
            log(f"      Aggregator: {len(agg_rows)} dòng · "
                f"{sum(r['sales'] for r in agg_rows):,.0f} đ".replace(",", "."))

    # ---- C4 · chi phí ngoài media ----
    try:
        g = _grid(hit, "Budget")
    except KeyError:
        g = None
    if g:
        h = _find_section(g, "C4.", "Hạng mục")
        items = []
        if h is not None:
            for r in g[h + 1:]:
                name = str(r[0]).strip() if r and r[0] else ""
                if not name:
                    continue
                if norm(name).startswith(("tổng", "tong", "c5")):
                    break
                b, a = to_num(r[1] if len(r) > 1 else None), to_num(r[2] if len(r) > 2 else None)
                if b is None and a is None:
                    continue
                items.append({"month": month, "item": name,
                              "budget": round(b) if b is not None else None,
                              "actual": round(a) if a is not None else None})
        if items:
            out["budget_nonmedia"] = items
            log(f"      Chi phí ngoài media: {len(items)} hạng mục")
    return out


# ══════════════════════════════════════════════════════════════════════
# 8. BOOKING TIỆC → booking
# ══════════════════════════════════════════════════════════════════════
def _pax(v):
    """'50 pax ' → 50 · '2' → 2 · '14 va 21/11' → 14 · trống/TBA → None."""
    m = re.search(r"\d+", str(v or ""))
    return int(m.group()) if m else None


_BOOKING_CACHE = {}


def _booking_leads(path):
    """Đọc sổ booking MỘT lần mỗi lượt chạy → list lead đã làm sạch (mọi tháng).

    read_booking() được gọi cho từng tháng; đọc lại file và in lại cảnh báo 16 lần
    vừa chậm vừa làm báo cáo cập nhật ngập chữ."""
    sig = (path, os.path.getmtime(path))
    if sig in _BOOKING_CACHE:
        return _BOOKING_CACHE[sig]
    from datetime import date as _date, datetime as _dt
    asof = _dt.fromtimestamp(os.path.getmtime(path)).strftime("%Y-%m-%d")

    rows = sheet_rows(path, "Sales Info", header_row=1)
    leads = []
    if rows:
        c0 = rows[0]
        k_id = col(c0, "lead id")
        k_guest = col(c0, exact="Guest")
        k_inq = col(c0, "inquiry date")
        k_in = col(c0, "input date")
        k_ev = col(c0, "date event (start)", "date event")
        k_out = col(c0, exact="Outlet")
        k_ty = col(c0, "event type")
        k_src = col(c0, exact="Source")
        k_st = col(c0, exact="Status")
        k_pax = col(c0, "number of guests")
        k_exp = col(c0, "expected revenue")
        k_cls = col(c0, "closed revenue")
        k_why = col(c0, "reason if lost")
        k_cmt = col(c0, exact="Comment")

        fixed, unmapped, dropped = [], set(), 0
        for r in rows:
            if not any(r.get(k) for k in (k_id, k_guest, k_st, k_src) if k):
                continue                                   # dòng trống cuối bảng
            ev = date_of(r.get(k_ev))
            if ev and ev < "2020":                          # '1900-01-01' = ô ngày gõ 0
                ev = None
            inq = date_of(r.get(k_inq)) or date_of(r.get(k_in))
            # BẪY: ngày nhận lead gõ đảo ngày/tháng ('2026-10-08' cho lead nhận 10/08,
            # sự kiện 26/08). Ngày nhận SAU ngày xuất file hoặc SAU ngày sự kiện là sai
            # chắc chắn → thử đảo; đảo vẫn vô lý thì bỏ, dùng ngày sự kiện.
            if inq and (inq > asof or (ev and inq > ev)):
                y, mo, d = int(inq[:4]), int(inq[5:7]), int(inq[8:10])
                swap = None
                if d <= 12:
                    try:
                        swap = _date(y, d, mo).strftime("%Y-%m-%d")
                    except ValueError:
                        swap = None
                ok = swap and swap <= asof and (not ev or swap <= ev)
                fixed.append(f"{r.get(k_id) or '?'}: {inq} → {swap if ok else 'ngày sự kiện'}")
                inq = swap if ok else None
            lead_day = inq or ev
            if not lead_day:
                dropped += 1
                continue

            outlet = str(r.get(k_out) or "").strip() or None
            store = (store_code(outlet) or store_in_text(outlet)) if outlet else None
            if outlet and not store and norm(outlet) not in ("tba", "tbd"):
                unmapped.add(outlet)
            status = str(r.get(k_st) or "").strip() or None
            stage = booking_stage(status)
            etype = booking_etype(r.get(k_ty))
            guests = _pax(r.get(k_pax))
            exp = to_num(r.get(k_exp))
            exp = exp if exp and exp > 0 else None          # trống / TBA / 0 = CHƯA báo giá
            closed = to_num(r.get(k_cls))
            closed = closed if stage == "won" and closed and closed > 0 else None
            leads.append({
                "month": lead_day[:7],
                "ev_month": ev[:7] if ev else None,
                "store": store,
                "brand": BRAND_OF_STORE.get(store) if store else None,
                "outlet": outlet,
                "seg": booking_segment(etype, guests),
                "etype": etype,
                "source": str(r.get(k_src) or "").strip() or None,
                "status": status,
                "stage": stage,
                "lost_reason": booking_lost_reason(r.get(k_why), r.get(k_cmt))
                               if stage == "lost" else None,
                "leads": 1,
                "guests": guests or 0,
                "exp": exp or 0,
                "exp_n": 1 if exp else 0,
                "closed": closed or 0,
                "closed_n": 1 if closed else 0,
                "inq_est": 0 if inq else 1,
            })

        n_tab = sum(1 for x in leads if x["seg"] == "table")
        n_est = sum(x["inq_est"] for x in leads)
        n_nocl = sum(1 for x in leads if x["stage"] == "won" and not x["closed_n"])
        log(f"      Booking (sổ luỹ kế): {len(leads)} lead · {n_tab} dòng đặt bàn nhỏ tách riêng · "
            f"{n_est} lead thiếu ngày nhận (dùng ngày sự kiện)")
        if fixed:
            warn(f"Booking: {len(fixed)} ngày nhận lead vô lý đã sửa — " + " · ".join(fixed[:5]))
        if n_nocl:
            warn(f"Booking: {n_nocl} lead Confirmed CHƯA nhập Closed Revenue — doanh thu chốt đang thiếu")
        if unmapped:
            warn("Booking: outlet chưa khai ở dim_store.aliases → " + ", ".join(sorted(unmapped)))
        if dropped:
            warn(f"Booking: {dropped} dòng không có ngày nhận lẫn ngày sự kiện — bị bỏ")
    _BOOKING_CACHE[sig] = leads
    return leads


def booking_months():
    """Mọi tháng NHẬN LEAD có trong sổ booking — update.py dựng đủ các tháng này."""
    path = l0_latest("S07_lead")
    return sorted({x["month"] for x in _booking_leads(path)}) if path else []


def read_booking(month):
    """Sổ booking → sheet `booking` của THÁNG NHẬN LEAD.

    Trước 17/09/2026 tháng của dòng là tháng DIỄN RA sự kiện. Sai cho phễu: chi phí
    ads của tháng 8 sinh ra lead tháng 8, nhưng tiệc có thể diễn ra tháng 11 — ghép
    theo tháng sự kiện là đem chi phí tháng này chia cho lead tháng khác. Nay `month`
    = tháng nhận lead (cohort), `ev_month` giữ tháng diễn ra cho lịch doanh thu."""
    path = l0_latest("S07_lead")
    if not path:
        warn("không thấy file Booking & Event")
        return {}
    sel = [x for x in _booking_leads(path) if x["month"] == month]
    if not sel:
        return {}
    keys = ["month", "ev_month", "store", "brand", "outlet", "seg", "etype", "source",
            "status", "stage", "lost_reason"]
    sums = ["leads", "guests", "exp", "exp_n", "closed", "closed_n", "inq_est"]
    out = agg(sel, keys, sums=sums)
    for r in out:
        for k in sums:
            r[k] = round(r[k])
    log(f"      Booking: {len(sel)} lead nhận trong tháng → {len(out)} dòng gộp · "
        f"chốt {sum(r['closed'] for r in out if r['stage'] == 'won'):,.0f} đ".replace(",", "."))
    return {"booking": sorted(out, key=lambda r: (-r["leads"], -r["exp"]))}


# ══════════════════════════════════════════════════════════════════════
# 9. PARTNERSHIP → partner_month
# ══════════════════════════════════════════════════════════════════════
def read_partnership(month):
    files = l0_month_files("S21_evoucher", month)
    if not files:
        return {}
    # File eVoucher chỉ là DANH SÁCH MÃ ĐÃ PHÁT — không có lượt dùng, doanh thu.
    # Vì vậy chỉ điền `issued`; used/rev/disc để TRỐNG (chưa đo được), không điền 0.
    issued = defaultdict(int)
    for f in sorted(files):
        from openpyxl import load_workbook
        wb = load_workbook(f, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        ws.reset_dimensions()
        n = sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0])
        wb.close()
        code = "P01" if "techcombank" in norm(f) else None
        if not code:
            continue
        issued[code] += n
        log(f"      Partnership: {os.path.basename(f)[:46]} → {n} mã")
    return {"partner_month": [{"month": month, "code": c, "issued": n}
                              for c, n in sorted(issued.items())]} if issued else {}



# ══════════════════════════════════════════════════════════════════════
# 10. POS THÔ (bảng kê hoá đơn + báo cáo bán hàng) → channel · daypart ·
#     identify · nature · recon
# ══════════════════════════════════════════════════════════════════════
# BẪY 1: dòng TỔNG nằm lẫn trong dữ liệu ở cả ba nguồn. Bảng kê ghi 'TẠI CHỖ' /
#        'MANG VỀ' ngay ở cột Mã hoá đơn. Nạp nhầm là nhân đôi toàn bộ.
TOTAL_MARKERS = {"tổng", "tong", "tổng:", "tại chỗ", "mang về", "grab", "corporate", ""}

# BẪY 2: từ T8/2026 iPOS xuất MỖI CỬA HÀNG MỘT SHEET rồi thêm sheet 'Tất cả cửa hàng'
#        ở cuối. Đọc sheet đầu tiên (như lane cũ vẫn làm) chỉ lấy được MỘT cửa hàng;
#        đọc hết mọi sheet thì cộng đôi. Quy tắc: có sheet tổng thì CHỈ đọc sheet tổng.
ALL_STORE_SHEET = "tất cả cửa hàng"

DAYPARTS = [(0, 10, "Sáng ≤10h"), (11, 14, "Trưa 11-14h"), (15, 17, "Chiều 15-17h"),
            (18, 21, "Tối 18-21h"), (22, 23, "Khuya 22h+")]
DAYPART_ORDER = [lb for _, _, lb in DAYPARTS]

# Bảng luật phân loại bản chất CTKM ĐÃ CHUYỂN về data_contract.json → $promo_nature,
# đọc qua monthly_lib.classify_nature. Trước đây file này giữ một bản sao riêng và nó
# đã lệch thật: bản ở đây bắt `nhân viên` trước `skg` nên nhân viên ĐỐI TÁC bị xếp
# vào nội bộ, và chưa có nhãn CARE nên CSKH bị tính là marketing.

def daypart_of(hour):
    if hour is None:
        return None
    for a, b, lb in DAYPARTS:
        if a <= hour <= b:
            return lb
    return "Khuya 22h+"


def _hour(v):
    m = re.match(r"^(\d{1,2})", str(v or "").strip())
    return int(m.group(1)) if m else None


def _clean(v):
    # BẪY 3: ô "rỗng" của iPOS là ký tự vô hình U+200B, không phải chuỗi rỗng.
    return str(v).replace("\u200b", "").replace("\ufeff", "").strip() if v is not None else ""


def _camp(v):
    """T\u00ean CTKM \u2014 gi\u1eef NGUY\u00caN V\u0102N (ch\u1eef hoa, d\u1ea5u) nh\u01b0ng g\u1ed9p kho\u1ea3ng tr\u1eafng th\u1eeba.

    `BUY 01 DRINK  GET 01 DRINK` (hai d\u1ea5u c\u00e1ch, lane d\u00f2ng m\u00f3n) v\u00e0
    `BUY 01 DRINK GET 01 DRINK` (m\u1ed9t d\u1ea5u c\u00e1ch, lane ho\u00e1 \u0111\u01a1n) l\u00e0 C\u00d9NG m\u1ed9t ch\u01b0\u01a1ng
    tr\u00ecnh g\u00f5 l\u1ec7ch tay. Kh\u00f4ng g\u1ed9p th\u00ec fact t\u00e1ch \u0111\u00f4i v\u00e0 m\u1ecdi ph\u00e9p c\u1ed9ng \u0111\u1ec1u h\u1ee5t."""
    return re.sub(r"\s+", " ", _clean(v))


def read_pos_sheets(path, want):
    """Đọc file POS thành list[dict] theo tên cột CHUẨN HOÁ của ta.

    `want` = {tên_chuẩn: [bí danh, …]} vì iPOS đổi tên cột giữa các kỳ export:
    T7 dùng `Thời gian` / `Mã hoá đơn`, T8 dùng `Ngày` / `Hoá đơn`. Dòng tiêu đề
    cũng lúc ở dòng 1, lúc ở dòng 2 (bên trên là dòng tựa đề) — nên dò theo nội
    dung chứ không đếm dòng."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    targets = [n for n in wb.sheetnames if norm(n) == ALL_STORE_SHEET] or list(wb.sheetnames)
    rows, missing = [], None
    for nm in targets:
        ws = wb[nm]
        ws.reset_dimensions()
        hdr = None
        for r in ws.iter_rows(values_only=True):
            if r is None:
                continue
            if hdr is None:
                cand = {norm(_clean(c)): j for j, c in enumerate(r) if c is not None}
                found = {}
                for key, aliases in want.items():
                    for a in aliases:
                        if norm(a) in cand:
                            found[key] = cand[norm(a)]
                            break
                if len(found) >= 2:
                    hdr = found
                    missing = [k for k in want if k not in found]
                continue
            rows.append({k: (r[j] if j < len(r) else None) for k, j in hdr.items()})
    wb.close()
    if missing:
        warn("%s: thiếu cột %s" % (os.path.basename(path)[:44], ", ".join(missing)))
    return rows, targets


def _pos_files(month):
    # Hai file cùng tháng (vd. bản 'tới 13-09' rồi bản đủ tháng) → lấy bản MỚI NHẤT.
    # Bản cũ dùng next() trên glob — thứ tự không xác định, có lúc lấy nhầm bản dở.
    return l0_month_file("S01_item", month), l0_month_file("S02_bill", month)


def read_pos(month, built=None):
    if not RAW:
        return {}
    item_f, bill_f = _pos_files(month)
    out = {}
    bills, bill_by_inv = [], {}

    # ---------- bảng kê hoá đơn ----------
    if bill_f:
        want = {
            "store": ["Cửa hàng"],
            "inv": ["Mã hoá đơn", "Hoá đơn"],
            "no": ["Số HĐ", "Số hoá đơn"],
            "guest": ["Số khách"],
            "channel": ["Nguồn"],
            "date": ["Ngày chứng từ", "Ngày", "Thời gian"],
            "hour": ["Giờ vào", "Giờ"],
            "net": ["Tổng tiền"],
            "phone": ["Số điện thoại"],
            # CTKM và chiết khấu CHỈ đọc được đúng ở đây. Báo cáo bán hàng cũng có
            # cột "Giảm giá" nhưng iPOS bỏ trống nó ở ~95% dòng vì chiết khấu được
            # ghi MỘT LẦN trên hoá đơn, không rải xuống từng món.
            "camp": ["Tên CTKM"],
            "gross": ["Thanh toán trước giảm giá"],
            "disc_amt": ["Giảm giá"],
            "disc_pct": ["Chiết khấu"],
            "disc_gg": ["Phiếu GG"],
            "net_novat": ["Tổng tiền (không bao gồm VAT)"],
        }
        rows, sheets = read_pos_sheets(bill_f, want)
        bills, seen, unknown = [], set(), set()
        bill_by_inv = {}              # mã hoá đơn → hoá đơn (để gắn dòng món LTO vào TỔNG hoá đơn)
        for r in rows:
            inv = _clean(r.get("inv"))
            if norm(inv) in TOTAL_MARKERS or not _clean(r.get("no")):
                continue
            code = store_code(r.get("store"))
            if not code:
                if _clean(r.get("store")):
                    unknown.add(_clean(r.get("store")))
                continue
            d = date_of(_clean(r.get("date")))
            if not d or d[:7] != month or inv in seen:
                continue
            seen.add(inv)
            bill_by_inv[inv] = len(bills)
            bills.append({
                "store": code, "tc": 1,
                "net": to_num(r.get("net"), 0) or 0,
                "guest": to_num(r.get("guest"), 0) or 0,
                "channel": _clean(r.get("channel")) or "(không rõ)",
                "daypart": daypart_of(_hour(r.get("hour"))),
                "phone": _clean(r.get("phone")),
                # khối CTKM cấp hoá đơn
                "date": d,
                "camp": _camp(r.get("camp")),
                "gross": to_num(r.get("gross"), 0) or 0,
                # `net` của khối CTKM = cột `Tổng tiền` (đã dùng ở trên), KHÔNG phải
                # `Tổng tiền (không bao gồm VAT)`. Đối soát T4/2026: tổng cột này
                # bằng ĐÚNG Net Sales của Tracking Sales (4.242.462.365) tới từng
                # đồng — tức đây mới là base mà daily/store_month đang dùng. Kỳ nền
                # của M7.2 lấy từ daily.json, nên hai vế phải cùng base thì lift mới
                # có nghĩa. `net_novat` giữ lại để đối soát, không xuất ra fact.
                "net_novat": to_num(r.get("net_novat"), 0) or 0,
                # GIẢM GIÁ ≠ PHIẾU GIẢM GIÁ — hai loại chi phí, để riêng.
                #
                # `disc` = `Giảm giá` (số tiền cố định) + `Chiết khấu` (theo %).
                #   Lane cũ chỉ lấy `Giảm giá` nên bỏ quên `Chiết khấu` — và
                #   `Chiết khấu` mới là phần lớn: T4/2026 là 372,4tr so với 13,4tr.
                #   Đó là toàn bộ lỗi “chi phí ưu đãi sai 20–45 lần”.
                #
                # `voucher` = `Phiếu GG`. Voucher là PHƯƠNG THỨC THANH TOÁN, không
                #   phải giảm giá — nên nó KHÔNG nằm trong `gross − net` của
                #   Tracking Sales. Gộp vào `disc` thì tổng chiết khấu vọt lên 147%
                #   chi phí ưu đãi toàn chuỗi, một con số không thể đúng. Tách ra
                #   thì T1 về 97%, T4 96%, T7 92% — hợp lý.
                "disc": (to_num(r.get("disc_amt"), 0) or 0)
                        + (to_num(r.get("disc_pct"), 0) or 0),
                "voucher": to_num(r.get("disc_gg"), 0) or 0,
            })
        if unknown:
            warn("bảng kê %s: cửa hàng chưa khai — %s" % (month, ", ".join(sorted(unknown))))
        if bills:
            out["channel"] = [
                {"month": month, "channel": r["channel"],
                 "net": round(r["net"]), "tc": round(r["tc"])}
                for r in sorted(agg(bills, ["channel"], sums=["net", "tc"]),
                                key=lambda x: -x["net"])]
            dp = agg([b for b in bills if b["daypart"]], ["daypart"],
                     sums=["net", "tc", "guest"])
            out["daypart"] = [
                {"month": month, "daypart": r["daypart"], "net": round(r["net"]),
                 "tc": round(r["tc"]), "guest": round(r["guest"])}
                for r in sorted(dp, key=lambda x: DAYPART_ORDER.index(x["daypart"]))]
            out["identify"] = [{"month": month, "bills": len(bills),
                                "id_bills": sum(1 for b in bills if b["phone"])}]
            out["_bills_n"] = len(bills)
            out["_bill_store"] = agg(bills, ["store"], sums=["net", "tc", "guest"])
            log("      POS bảng kê: %s hoá đơn · %d sheet"
                % (format(len(bills), ",").replace(",", "."), len(sheets)))

            # ---------- nền của M7.2: CTKM × ngày × cửa hàng ----------
            ck = [b for b in bills if b["camp"] and b["date"]]
            if ck:
                for b in ck:
                    b["brand"] = BRAND_OF_STORE.get(b["store"], "OTHER")
                    b["nature"] = classify_nature(b["camp"])
                    b["name_pos"] = b["camp"]     # giữ NGUYÊN VĂN, chuẩn hoá ở dim_campaign
                    b["guests"] = b["guest"]
                    b["bills"] = 1
                out["_promo_bill"] = agg(
                    ck, ["date", "store", "name_pos"],
                    sums=["bills", "guests", "gross", "disc", "voucher", "net"],
                    first=["brand", "nature"])
                out["_promo_disc"] = agg(ck, ["nature", "brand"], sums=["disc", "voucher"])
                log("      CTKM cấp hoá đơn: %s bill · chiết khấu %s"
                    % (format(len(ck), ",").replace(",", "."),
                       format(round(sum(b["disc"] for b in ck)), ",").replace(",", ".")))
    else:
        warn("không thấy bảng kê hoá đơn của %s" % month)

    # ---------- báo cáo bán hàng (cấp món) ----------
    if item_f:
        want = {
            "store": ["Cửa hàng"],
            "date": ["Ngày", "Thời gian"],
            "net": ["Tổng tiền"],
            "line_rev": ["Thành tiền"],
            "camp": ["Tên CTKM"],
            "disc": ["Giảm giá"],
            "inv": ["Mã hoá đơn", "Hoá đơn"],
            "code": ["Mã hàng"],
            "name": ["Tên hàng"],
            "group": ["Nhóm món"],
            "qty": ["Số lượng"],
        }
        rows, sheets = read_pos_sheets(item_f, want)
        lto_rx = re.compile(CONTRACT["$campaign"].get("lto_group_rx", "LTO"), re.I)
        lto_lines = []
        items = []
        for r in rows:
            code = store_code(r.get("store"))
            if not code:
                continue
            d = date_of(_clean(r.get("date")))
            if not d or d[:7] != month:
                continue
            grp = _clean(r.get("group"))
            if grp and lto_rx.search(grp):
                lto_lines.append({"date": d, "store": code, "inv": _clean(r.get("inv")),
                                  "item_code": _clean(r.get("code")), "item_name": _clean(r.get("name")),
                                  "group": grp, "qty": to_num(r.get("qty"), 0) or 0,
                                  "line_rev": to_num(r.get("line_rev"), 0) or 0})
            items.append({
                "store": code, "date": d, "items": 1,
                "net": to_num(r.get("net"), 0) or 0,           # tổng tiền dòng món
                # `rev` của khối CTKM tính trên THÀNH TIỀN (trước phí dịch vụ và VAT) —
                # đây là quy ước đã chốt cho cơ cấu chương trình, khác với Net Sales.
                "rev": to_num(r.get("line_rev"), 0) or 0,
                # Cột `Giảm giá` CẤP DÒNG MÓN — iPOS bỏ trống ~95% vì chiết khấu ghi
                # ở cấp hoá đơn. Giữ lại CHỈ để làm phương án dự phòng khi không có
                # bảng kê; số thật lấy từ `_promo_disc`. Đối soát T1–T7/2026: cột này
                # cho 8–18tr/tháng trong khi chiết khấu thật là 218–386tr — lệch 18–45 lần.
                "disc": to_num(r.get("disc"), 0) or 0,
                "camp": _camp(r.get("camp")), "inv": _clean(r.get("inv"))})
        if items:
            ck = [i for i in items if i["camp"]]
            if ck:
                for i in ck:
                    i["nature"] = classify_nature(i["camp"])
                    i["brand"] = BRAND_OF_STORE.get(i["store"], "OTHER")
                    i["name_pos"] = i["camp"]
                out["_promo_item"] = agg(ck, ["date", "store", "name_pos"],
                                         sums=["rev", "items"])
                # Chiết khấu thật đo ở cấp hoá đơn. Chỉ rơi về cột cấp dòng món khi
                # tháng đó KHÔNG có bảng kê — và khi ấy phải nói ra là số bị hụt.
                # `nature.disc` = CHI PHÍ ƯU ĐÃI TỔNG (giảm giá + phiếu GG) — đúng với
                # tên gọi trên màn hình M7. M7.2 cần tách hai loại thì đọc
                # fact_promo_day, ở đó `disc` và `voucher` nằm riêng.
                disc_by = {(r["nature"], r["brand"]): r["disc"] + r["voucher"]
                           for r in out.get("_promo_disc", [])}
                if not disc_by:
                    warn("%s: không có bảng kê hoá đơn — chi phí ưu đãi lấy tạm ở cấp "
                         "dòng món, HỤT khoảng 20–45 lần, đừng dùng để tính ROI" % month)
                    disc_by = {(r["nature"], r["brand"]): r["disc"]
                               for r in agg(ck, ["nature", "brand"], sums=["disc"])}
                bills_by = defaultdict(set)
                for i in ck:
                    bills_by[(i["nature"], i["brand"])].add(i["inv"])
                out["nature"] = [
                    {"month": month, "nature": r["nature"], "brand": r["brand"],
                     "rev": round(r["rev"]),
                     "disc": round(disc_by.get((r["nature"], r["brand"]), 0)),
                     "bills": len(bills_by[(r["nature"], r["brand"])])}
                    for r in sorted(agg(ck, ["nature", "brand"], sums=["rev", "disc"]),
                                    key=lambda x: (x["nature"], x["brand"]))]
            out["_item_store"] = agg(items, ["store"], sums=["net"])
            # Số dòng món của tháng đi kèm sheet identify để tab D1 đếm được
            # tổng dòng đã xử lý mà không cần khai cứng ở _stats.
            if out.get("identify"):
                out["identify"][0]["items"] = len(items)
            log("      POS bán hàng: %s dòng món · %d sheet"
                % (format(len(items), ",").replace(",", "."), len(sheets)))
        # ---------- fact_lto_line — chương trình LTO chạy theo MÓN, không có tên CTKM trên hoá đơn ----------
        # Mỗi dòng món thuộc nhóm LTO (Nhóm món chứa "LTO": LTO DRINK · LTO MENU · NESTLE LTO) gắn với
        # TỔNG hoá đơn chứa nó. Doanh thu CTKM của chương trình LTO = Σ Tổng tiền các hoá đơn có món LTO
        # (gồm cả món khác khách gọi kèm); doanh thu riêng món LTO = Σ Thành tiền dòng LTO.
        if lto_lines:
            bl = bills
            miss = 0
            rows_lto = []
            for ln in lto_lines:
                bi = bill_by_inv.get(ln["inv"])
                b = bl[bi] if bi is not None and bi < len(bl) else None
                if b is None:
                    miss += 1
                code_ = ln["item_code"] or ""
                rows_lto.append({
                    "date": ln["date"], "store": ln["store"], "brand": BRAND_OF_STORE.get(ln["store"], "OTHER"),
                    "bill": ln["inv"], "item_code": code_,
                    # mã gốc: bỏ hậu tố MK (cùng món, khác bảng giá) — LTD0009MK ≡ LTD0009
                    "item_base": re.sub(r"MK$", "", code_.upper()),
                    "item_name": ln["item_name"], "group": ln["group"],
                    "qty": round(ln["qty"], 2), "line_rev": round(ln["line_rev"]),
                    "bill_net": round(b["net"]) if b else None, "bill_gross": round(b["gross"]) if b else None,
                    "bill_disc": round(b["disc"]) if b else None, "bill_voucher": round(b["voucher"]) if b else None,
                    "bill_guests": round(b["guest"]) if b else None, "bill_camp": (b or {}).get("camp"),
                })
            out["fact_lto_line"] = rows_lto
            log("      món LTO: %d dòng · %d hoá đơn%s" % (len(rows_lto), len({r["bill"] for r in rows_lto}),
                                                         " · %d dòng không khớp bảng kê" % miss if miss else ""))
    else:
        warn("không thấy báo cáo bán hàng của %s" % month)

    # ---------- fact_promo_day — nền của M7.2 ----------
    # Hai lane đo hai thứ khác nhau trên cùng một chương trình và KHÔNG được cộng
    # chung: bảng kê cho số bill, số khách và chiết khấu THẬT (cấp hoá đơn); báo
    # cáo bán hàng cho doanh thu chạm theo THÀNH TIỀN (cấp dòng món). Giữ cả hai,
    # đặt cạnh nhau, khoá chung là ngày + cửa hàng + tên CTKM nguyên văn.
    pb = {(r["date"], r["store"], r["name_pos"]): r for r in out.pop("_promo_bill", [])}
    pi = {(r["date"], r["store"], r["name_pos"]): r for r in out.pop("_promo_item", [])}
    out.pop("_promo_disc", None)
    if pb or pi:
        rows = []
        for k in sorted(set(pb) | set(pi)):
            b, i = pb.get(k), pi.get(k)
            d, store, name = k
            rows.append({
                "date": d, "store": store,
                "brand": (b or {}).get("brand") or BRAND_OF_STORE.get(store, "OTHER"),
                "name_pos": name,
                "nature": (b or {}).get("nature") or classify_nature(name),
                "bills": round(b["bills"]) if b else None,
                "guests": round(b["guests"]) if b else None,
                "gross": round(b["gross"]) if b else None,
                "disc": round(b["disc"]) if b else None,
                "voucher": round(b["voucher"]) if b else None,
                "net": round(b["net"]) if b else None,
                "rev": round(i["rev"]) if i else None,
                "items": round(i["items"]) if i else None,
            })
        out["fact_promo_day"] = rows
        only_item = len(set(pi) - set(pb))
        if only_item:
            warn("%s: %d dòng CTKM chỉ thấy ở báo cáo bán hàng, không thấy ở bảng kê "
                 "— chưa có chiết khấu" % (month, only_item))

    # ---------- đối soát ba tầng ----------
    out.pop("_bills_n", None)
    sm = {r["store"]: r for r in (built or {}).get("store_month", [])}
    bl = {r["store"]: r for r in out.pop("_bill_store", [])}
    it = {r["store"]: r for r in out.pop("_item_store", [])}
    if sm and (bl or it):
        out["recon"] = [{
            "month": month, "store": c,
            "net": sm[c]["net"], "tc": sm[c]["tc"], "guest": sm[c]["guest"],
            "net_item": round(it[c]["net"]) if c in it else None,
            "net_bill": round(bl[c]["net"]) if c in bl else None,
            "tc_bill": round(bl[c]["tc"]) if c in bl else None,
            "guest_bill": round(bl[c]["guest"]) if c in bl else None,
        } for c in sorted(sm)]
    return out

# ══════════════════════════════════════════════════════════════════════
# CHẠY
# ══════════════════════════════════════════════════════════════════════
# (mã phần, nhãn, hàm, nguồn L0 nuôi phần đó). update.py dựa vào cột nguồn để chỉ
# chạy lại PHẦN bị ảnh hưởng: thả file Meta Ads T8 thì chỉ dựng lại phần meta của
# T8, không đọc lại 60MB file POS.
BUILDERS = [
    ("tracking", "Doanh thu · Tracking Sales", read_tracking, ["S03_daily", "S00_targets"]),
    ("pos", "POS · kênh · giờ · CTKM", read_pos, ["S01_item", "S02_bill"]),
    ("meta", "Meta Ads", read_meta_ads, ["S08_ads_meta"]),
    ("google", "Google Ads", read_google_ads, ["S09_ads_google"]),
    ("social", "Social · Facebook + TikTok", read_social, ["S18_social", "S22_tiktok"]),
    ("oa", "Zalo OA", read_oa, ["S12_zalo_oa"]),
    ("member", "Member đăng ký", read_member, ["S13_member"]),
    ("promotion", "Promotion & Aggregator", read_promotion, ["S19_aggregator"]),
    ("booking", "Booking tiệc", read_booking, ["S07_lead"]),
    ("partnership", "Partnership", read_partnership, ["S21_evoucher"]),
]
PART_SOURCES = {k: srcs for k, _, _, srcs in BUILDERS}
# Sheet mà MỖI phần sở hữu trong file tháng. Phần đã chạy xong mà không ra số → sheet
# của nó bị GỠ khỏi file tháng. Bản trước giữ nguyên sheet cũ: gỡ hoặc sửa file nguồn
# thì số cũ vẫn nằm đó mãi — đúng kiểu "số đóng băng" đã làm 5 nguồn chết âm thầm.
PART_SHEETS = {
    "tracking": ["store_month", "daily", "coverage", "dim_target"],
    "pos": ["channel", "daypart", "identify", "nature", "fact_promo_day", "fact_lto_line", "recon"],
    "meta": ["ads_month", "ads_brand", "ads_objective", "ads_campaign_detail"],
    "google": ["ads_google", "gads_channel", "gads_kw"],
    "social": ["social_month"],
    "oa": ["oa"],
    "member": ["member"],
    "promotion": ["aggregator", "budget_nonmedia"],
    "booking": ["booking"],
    "partnership": ["partner_month"],
}


def build(month, dry=False, parts=None):
    log(f"\n▸ {month}" + (f"  (chỉ dựng: {', '.join(parts)})" if parts else ""))
    path = os.path.join(MONTHLY_DIR, f"{month}.xlsx")
    keep = read_workbook(path) if os.path.exists(path) else {}
    built, owned = {}, set()
    for key, label, fn, _ in BUILDERS:
        if parts and key not in parts:
            continue
        try:
            # read_pos cần store_month để lập bảng đối soát — lấy bản vừa dựng, không
            # có thì lấy bản đang nằm trong file tháng.
            ctx = {"store_month": built.get("store_month") or keep.get("store_month", [])}
            got = (fn(month, ctx) if fn is read_pos else fn(month)) or {}
        except Exception as e:                                      # noqa: BLE001
            # LỖI thì giữ số cũ (không sở hữu sheet) — khác hẳn "chạy xong, không có số".
            warn(f"{label}: {type(e).__name__} — {str(e)[:90]}")
            got = None
        if got is not None:
            owned.update(PART_SHEETS.get(key, []))
        if got:
            built.update(got)
            log(f"      ✔ {label:<26} {' · '.join(f'{k}({len(v)})' for k, v in got.items())}")
        else:
            log(f"      – {label:<26} chưa có dữ liệu")

    dropped = [k for k in keep if k in owned and k not in built]
    if dropped:
        log(f"      ↳ gỡ sheet không còn nguồn: {', '.join(dropped)}")
    merged = {k: v for k, v in keep.items() if k not in dropped}
    merged.update(built)                       # sheet dựng được thì ghi đè
    order = [k for k in merged if k in built] + [k for k in merged if k not in built]
    merged = {k: merged[k] for k in dict.fromkeys(order)}

    if dry:
        log(f"      (dry) sẽ ghi {len(merged)} sheet vào {os.path.relpath(path, ROOT)}")
        return
    write_workbook(path, merged,
                   title=f"{month} — SỐ LIỆU THÁNG (một tháng một file)",
                   guide_lines=guide_monthly(month))
    n = sum(len(v) for v in merged.values())
    log(f"      → {os.path.relpath(path, ROOT)}  {len(merged)} sheet · {n:,} dòng"
        .replace(",", "."))


def months_available():
    """Tháng nào có mặt ở nguồn doanh thu — dùng cho --all."""
    if not TRACKING or not os.path.exists(TRACKING):
        return []
    ms = set()
    for r in sheet_rows(TRACKING, "Data_Daily", header_row=4):
        d = date_of(r.get("Ngày"))
        if d:
            ms.add(d[:7])
    return sorted(ms)


def main(argv):
    dry = "--dry" in argv
    parts = None
    for a in argv:
        if a.startswith("--parts="):
            parts = [x for x in a.split("=", 1)[1].split(",") if x]
            bad = [x for x in parts if x not in PART_SOURCES]
            if bad:
                log(f"phần không tồn tại: {bad} · hợp lệ: {', '.join(PART_SOURCES)}")
                return 1
    args = [a for a in argv if not a.startswith("--")]
    log("─" * 74)
    log("NOIRE — dựng gói số liệu tháng từ dữ liệu thô")
    log(f"gốc dữ liệu: {RAW_ROOT or '(KHÔNG TÌM THẤY — đặt biến NOIRE_ROOT)'}")
    log(f"chọn vì   : {L0_WHY}")
    log("─" * 74)
    if not RAW_ROOT:
        return 1

    months = months_available() if "--all" in argv else [m for m in args if month_of(m)]
    if not months:
        log("Cách dùng: python tools/build_month.py 2026-08 [2026-07 …] [--all] [--dry]")
        return 1

    os.makedirs(MONTHLY_DIR, exist_ok=True)
    for m in months:
        build(m, dry, parts)

    log("\n" + "─" * 74)
    if WARN:
        log(f"{len(WARN)} cảnh báo:")
        for w in WARN:
            log(f"   ! {w}")
    else:
        log("Không có cảnh báo.")
    log("Bước tiếp theo:  npm run build:data")
    log("─" * 74)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
