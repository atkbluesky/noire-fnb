"""Sinh mẫu ngân sách Marketing chuẩn (cửa hàng × brand × tháng × quý × kênh ads).

    python tools/make_budget_template.py            -> Q4/2026
    python tools/make_budget_template.py 2027 1     -> Q1/2027

Ghi ra L0_input/03_MARKETING/04_Ngan_Sach/_MAU_NGAN_SACH_Q{q}_{year}.xlsx
Người dùng chỉ nhập số vào các ô vàng của sheet NHAP_NGAN_SACH; sheet TONG_HOP tự tính.
"""
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

YEAR = int(sys.argv[1]) if len(sys.argv) > 1 else 2026
Q = int(sys.argv[2]) if len(sys.argv) > 2 else 4
MONTHS = [f"{YEAR}-{m:02d}" for m in range(3 * Q - 2, 3 * Q + 1)]
MLABEL = [f"Tháng {m}" for m in range(3 * Q - 2, 3 * Q + 1)]

# Danh mục cửa hàng lấy từ ngân sách Q3/2026 đã duyệt (thứ tự giữ nguyên)
STORES = [("NCB", "NOIRE ET"), ("NCB", "NOIRE The Mett"), ("NCB", "NOIRE SKC"),
          ("NDC", "NOIRE 39 NTMK"), ("NDC", "NOIRE Berkley"),
          ("NJFB", "NOIRE JFB CREST"), ("NJFB", "NOIRE JFB SSV")]
BRANDS = ["NCB", "NDC", "NJFB"]
CHANNELS = ["Google Ads", "Meta Ads"]
SPARE = 8  # dòng trống để thêm cửa hàng/kênh mới

INPUT = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="D9E1F2")
TOT = PatternFill("solid", fgColor="E2EFDA")
GREY = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
NUM = '#,##0;[Red]-#,##0;"-"'
PCT = "0.0%"


def head(ws, row, values, col0=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col0 + i, value=v)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def put(ws, ref, value, fmt=None, fill=None, bold=False):
    c = ws[ref]
    c.value = value
    c.border = BOX
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if bold: c.font = Font(bold=True)
    return c


wb = Workbook()

# ───────────────────────── HUONG_DAN
g = wb.active
g.title = "HUONG_DAN"
lines = [
    (f"NGÂN SÁCH MARKETING · Q{Q}/{YEAR}", True),
    ("Mẫu chuẩn — ngân sách theo CỬA HÀNG × BRAND × THÁNG × QUÝ, chia theo KÊNH digital ads (Google Ads, Meta Ads).", False),
    ("", False),
    ("CÁCH DÙNG", True),
    ("1. Sang sheet NHAP_NGAN_SACH, nhập số tiền (VND) vào các ô VÀNG: mỗi dòng = 1 cửa hàng × 1 kênh, 3 cột = 3 tháng của quý.", False),
    ("2. Cột Tổng quý, sheet TONG_HOP và các dòng kiểm tra tự tính — KHÔNG gõ đè.", False),
    ("3. Khoản không chi thì nhập 0 (để trống = chưa nhập; hệ thống sẽ báo thiếu).", False),
    ("4. Cửa hàng/kênh mới: dùng dòng trống cuối bảng, chọn Brand và Kênh từ danh sách thả xuống.", False),
    ("5. Lưu file, giữ tên dạng NOIRE_MKT_Q{q}_{năm}_Checked_....xlsx rồi thả vào L0_input/03_MARKETING/04_Ngan_Sach/.", False),
    ("", False),
    ("QUY ƯỚC — không đổi để hệ thống đọc được", True),
    ("• Hàng tiêu đề của bảng nhập là hàng 6; hàng 5 chứa mã tháng dạng YYYY-MM (đừng sửa).", False),
    ("• Cột A = Brand (NCB / NDC / NJFB) · B = Cửa hàng (đúng tên trong dim_store) · C = Kênh.", False),
    ("• Đơn vị: VND, số nguyên, chưa gồm VAT trừ khi ghi chú khác.", False),
    ("• Tên cửa hàng phải khớp danh mục cửa hàng của hệ thống; sai tên → dòng đó bị báo 'chưa map'.", False),
    ("", False),
    ("TẠO MẪU CHO QUÝ SAU:  python tools/make_budget_template.py <năm> <quý>", False),
]
for i, (t, b) in enumerate(lines, start=1):
    c = g.cell(row=i, column=1, value=t.replace("{q}", str(Q)).replace("{năm}", str(YEAR)))
    c.font = Font(bold=b, size=14 if i == 1 else 11)
g.column_dimensions["A"].width = 130

# ───────────────────────── NHAP_NGAN_SACH
n = wb.create_sheet("NHAP_NGAN_SACH")
n["A1"] = f"NHẬP NGÂN SÁCH ADS THEO CỬA HÀNG · Q{Q}/{YEAR}"
n["A1"].font = Font(bold=True, size=14)
n["A2"] = "Ô vàng = nhập tay (VND). Mỗi dòng: 1 cửa hàng × 1 kênh × 3 tháng."
n["A3"] = "Trạng thái nhập"
n["A3"].font = Font(bold=True)
FIRST = 7
nrows = len(STORES) * len(CHANNELS)
LAST_FIXED = FIRST + nrows - 1
LAST = LAST_FIXED + SPARE
n["B3"] = (f'=IF(COUNT(D{FIRST}:F{LAST_FIXED})={nrows * 3},"Đã nhập đủ",'
           f'"Còn "&({nrows * 3}-COUNT(D{FIRST}:F{LAST_FIXED}))&" ô cần nhập")')
n["B3"].font = Font(bold=True, color="C00000")
# hàng 5: mã tháng cho máy; hàng 6: tiêu đề cho người
for j, m in enumerate(MONTHS):
    c = n.cell(row=5, column=4 + j, value=m)
    c.font = Font(italic=True, color="808080", size=9)
    c.alignment = Alignment(horizontal="center")
n["G5"] = f"Q{Q}-{YEAR}"
n["G5"].font = Font(italic=True, color="808080", size=9)
n["G5"].alignment = Alignment(horizontal="center")
head(n, 6, ["Brand", "Cửa hàng", "Kênh"] + MLABEL + [f"Tổng Q{Q}", "Ghi chú"])

r = FIRST
for brand, store in STORES:
    for ch in CHANNELS:
        put(n, f"A{r}", brand, fill=GREY)
        put(n, f"B{r}", store, fill=GREY)
        put(n, f"C{r}", ch, fill=GREY)
        for col in "DEF":
            put(n, f"{col}{r}", None, NUM, INPUT)
        put(n, f"G{r}", f'=IF(COUNT(D{r}:F{r})=0,"",SUM(D{r}:F{r}))', NUM, SUB)
        put(n, f"H{r}", None)
        r += 1
for _ in range(SPARE):
    for col in "ABC":
        put(n, f"{col}{r}", None, fill=INPUT)
    for col in "DEF":
        put(n, f"{col}{r}", None, NUM, INPUT)
    put(n, f"G{r}", f'=IF(COUNT(D{r}:F{r})=0,"",SUM(D{r}:F{r}))', NUM, SUB)
    put(n, f"H{r}", None)
    r += 1
TOTROW = r
put(n, f"A{TOTROW}", "TỔNG", fill=TOT, bold=True)
put(n, f"B{TOTROW}", None, fill=TOT)
put(n, f"C{TOTROW}", None, fill=TOT)
for col in "DEFG":
    put(n, f"{col}{TOTROW}", f"=SUM({col}{FIRST}:{col}{LAST})", NUM, TOT, True)
put(n, f"H{TOTROW}", None, fill=TOT)

# kiểm tra dòng: có số mà thiếu Brand/Cửa hàng/Kênh
n[f"A{TOTROW + 2}"] = "Kiểm tra: dòng có số nhưng thiếu Brand/Cửa hàng/Kênh"
n[f"D{TOTROW + 2}"] = (f'=SUMPRODUCT((D{FIRST}:D{LAST}+E{FIRST}:E{LAST}+F{FIRST}:F{LAST}<>0)'
                       f'*((A{FIRST}:A{LAST}="")+(B{FIRST}:B{LAST}="")+(C{FIRST}:C{LAST}="")>0))')
n[f"E{TOTROW + 2}"] = f'=IF(D{TOTROW + 2}=0,"OK","LỖI: bổ sung Brand/Cửa hàng/Kênh")'
n[f"A{TOTROW + 3}"] = "Kiểm tra: mỗi cửa hàng chỉ thuộc 1 brand"
n[f"D{TOTROW + 3}"] = (f'=SUMPRODUCT((B{FIRST}:B{LAST}<>"")*(COUNTIFS(B{FIRST}:B{LAST},B{FIRST}:B{LAST},A{FIRST}:A{LAST},"<>"&A{FIRST}:A{LAST})>0))')
n[f"E{TOTROW + 3}"] = f'=IF(D{TOTROW + 3}=0,"OK","LỖI: cùng cửa hàng nhưng khác brand")'

dv_b = DataValidation(type="list", formula1='"' + ",".join(BRANDS) + '"', allow_blank=True)
dv_c = DataValidation(type="list", formula1='"' + ",".join(CHANNELS) + '"', allow_blank=True)
dv_n = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                      errorTitle="Sai giá trị", error="Nhập số tiền VND ≥ 0")
for dv in (dv_b, dv_c, dv_n):
    n.add_data_validation(dv)
dv_b.add(f"A{FIRST}:A{LAST}")
dv_c.add(f"C{FIRST}:C{LAST}")
dv_n.add(f"D{FIRST}:F{LAST}")
for col, w in zip("ABCDEFGH", (9, 22, 13, 17, 17, 17, 19, 34)):
    n.column_dimensions[col].width = w
n.freeze_panes = "D7"

# ───────────────────────── TONG_HOP
t = wb.create_sheet("TONG_HOP")
t["A1"] = f"TỔNG HỢP NGÂN SÁCH Q{Q}/{YEAR} — tự tính từ NHAP_NGAN_SACH"
t["A1"].font = Font(bold=True, size=14)
NS = "NHAP_NGAN_SACH"
A = f"{NS}!$A${FIRST}:$A${LAST}"
B = f"{NS}!$B${FIRST}:$B${LAST}"
C = f"{NS}!$C${FIRST}:$C${LAST}"
MC = [f"{NS}!${c}${FIRST}:${c}${LAST}" for c in "DEF"]


def sumifs(month_idx, *conds):
    parts = "".join(f",{rng},{crit}" for rng, crit in conds)
    return f"SUMIFS({MC[month_idx]}{parts})"


# 1. theo Brand
row = 3
t[f"A{row}"] = "1. Theo BRAND × tháng"
t[f"A{row}"].font = Font(bold=True)
row += 1
head(t, row, ["Brand"] + MLABEL + [f"Tổng Q{Q}", "% tổng chuỗi", "Ngân sách được duyệt (nhập)", "Chênh lệch"])
row += 1
b0 = row
for b in BRANDS:
    put(t, f"A{row}", b, bold=True)
    for j in range(3):
        put(t, f"{'BCD'[j]}{row}", "=" + sumifs(j, (A, f"$A{row}")), NUM)
    put(t, f"E{row}", f"=SUM(B{row}:D{row})", NUM, SUB)
    put(t, f"F{row}", f"=IF($E${b0 + 3}=0,\"\",E{row}/$E${b0 + 3})", PCT)
    put(t, f"G{row}", None, NUM, INPUT)
    put(t, f"H{row}", f'=IF(G{row}="","",E{row}-G{row})', NUM)
    row += 1
put(t, f"A{row}", "TỔNG", fill=TOT, bold=True)
for col in "BCDE":
    put(t, f"{col}{row}", f"=SUM({col}{b0}:{col}{row - 1})", NUM, TOT, True)
put(t, f"F{row}", None, fill=TOT)
put(t, f"G{row}", f'=IF(COUNT(G{b0}:G{row - 1})=0,"",SUM(G{b0}:G{row - 1}))', NUM, TOT, True)
put(t, f"H{row}", f'=IF(G{row}="","",E{row}-G{row})', NUM, TOT, True)
grand = row

# 2. theo Brand × Kênh
row += 3
t[f"A{row}"] = "2. Theo BRAND × KÊNH × tháng"
t[f"A{row}"].font = Font(bold=True)
row += 1
head(t, row, ["Brand", "Kênh"] + MLABEL + [f"Tổng Q{Q}"])
row += 1
for b in BRANDS:
    for ch in CHANNELS:
        put(t, f"A{row}", b, bold=True)
        put(t, f"B{row}", ch)
        for j in range(3):
            put(t, f"{'CDE'[j]}{row}", "=" + sumifs(j, (A, f"$A{row}"), (C, f"$B{row}")), NUM)
        put(t, f"F{row}", f"=SUM(C{row}:E{row})", NUM, SUB)
        row += 1

# 3. theo Cửa hàng
row += 2
t[f"A{row}"] = "3. Theo CỬA HÀNG × tháng (Google + Meta)"
t[f"A{row}"].font = Font(bold=True)
row += 1
head(t, row, ["Cửa hàng", "Brand"] + MLABEL + [f"Tổng Q{Q}", "Google Ads", "Meta Ads", "% Google"])
row += 1
s0 = row
for brand, store in STORES:
    put(t, f"A{row}", store, bold=True)
    put(t, f"B{row}", brand)
    for j in range(3):
        put(t, f"{'CDE'[j]}{row}", "=" + sumifs(j, (B, f"$A{row}")), NUM)
    put(t, f"F{row}", f"=SUM(C{row}:E{row})", NUM, SUB)
    put(t, f"G{row}", f'=SUMIFS({NS}!$G${FIRST}:$G${LAST},{B},$A{row},{C},"Google Ads")', NUM)
    put(t, f"H{row}", f'=SUMIFS({NS}!$G${FIRST}:$G${LAST},{B},$A{row},{C},"Meta Ads")', NUM)
    put(t, f"I{row}", f'=IF(F{row}=0,"",G{row}/F{row})', PCT)
    row += 1
put(t, f"A{row}", "TỔNG", fill=TOT, bold=True)
put(t, f"B{row}", None, fill=TOT)
for col in "CDEFGH":
    put(t, f"{col}{row}", f"=SUM({col}{s0}:{col}{row - 1})", NUM, TOT, True)
put(t, f"I{row}", f'=IF(F{row}=0,"",G{row}/F{row})', PCT, TOT, True)

# 4. theo Kênh
row += 3
t[f"A{row}"] = "4. Theo KÊNH × tháng"
t[f"A{row}"].font = Font(bold=True)
row += 1
head(t, row, ["Kênh"] + MLABEL + [f"Tổng Q{Q}", "% mix"])
row += 1
k0 = row
for ch in CHANNELS:
    put(t, f"A{row}", ch, bold=True)
    for j in range(3):
        put(t, f"{'BCD'[j]}{row}", "=" + sumifs(j, (C, f"$A{row}")), NUM)
    put(t, f"E{row}", f"=SUM(B{row}:D{row})", NUM, SUB)
    put(t, f"F{row}", f'=IF($E${k0 + len(CHANNELS)}=0,"",E{row}/$E${k0 + len(CHANNELS)})', PCT)
    row += 1
put(t, f"A{row}", "TỔNG", fill=TOT, bold=True)
for col in "BCDE":
    put(t, f"{col}{row}", f"=SUM({col}{k0}:{col}{row - 1})", NUM, TOT, True)
put(t, f"F{row}", None, fill=TOT)

# 5. đối chiếu
row += 3
t[f"A{row}"] = "5. Đối chiếu (phải bằng 0)"
t[f"A{row}"].font = Font(bold=True)
row += 1
put(t, f"A{row}", "Bảng nhập − Theo Brand")
put(t, f"E{row}", f"={NS}!G{TOTROW}-E{grand}", NUM)
put(t, f"F{row}", f'=IF(ROUND(E{row},0)=0,"OK","LỆCH: kiểm tra Brand ở bảng nhập")')
row += 1
put(t, f"A{row}", "Bảng nhập − Theo Cửa hàng")
put(t, f"E{row}", f"={NS}!G{TOTROW}-F{s0 + len(STORES)}", NUM)
put(t, f"F{row}", f'=IF(ROUND(E{row},0)=0,"OK","LỆCH: tên cửa hàng ngoài danh mục ở mục 3")')
row += 1
put(t, f"A{row}", "Bảng nhập − Theo Kênh")
put(t, f"E{row}", f"={NS}!G{TOTROW}-E{k0 + len(CHANNELS)}", NUM)
put(t, f"F{row}", f'=IF(ROUND(E{row},0)=0,"OK","LỆCH: kênh ngoài Google/Meta")')

t.column_dimensions["A"].width = 24
for col in "BCDEFGHI":
    t.column_dimensions[col].width = 19

out = Path(__file__).resolve().parent.parent / "L0_input" / "03_MARKETING" / "04_Ngan_Sach" / f"_MAU_NGAN_SACH_Q{Q}_{YEAR}.xlsx"
wb.save(out)
print("đã ghi", out)
