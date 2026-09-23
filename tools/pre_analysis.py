# -*- coding: utf-8 -*-
"""
ĐỌC FILE PRE-ANALYSIS (S16) — BỘ ĐỌC DUY NHẤT
=============================================
    L0_input/03_MARKETING/05_Promotion_Ke_Hoach/NOIRE_Promotion_Pre-Analysis_*.xlsx

M7.1 Pre-Analytics (kế hoạch) và M7.2 Promotion Tracking (target để chấm ĐẠT) cùng đọc
qua hàm này — không nơi nào chép tay số kế hoạch sang file khác.

Cấu trúc file (theo slide TC_AOV_FnB_Marketing.pdf · bước 1–3 slide 8):
  · 4 sheet loại có bảng chuẩn, dòng tiêu đề chứa "ID" + "Chương trình":
        Combo-SetMenu (C#) · Discount (D#) · Gift-FOC (G#) · LTO-Item (L#)
  · Voucher-Loyalty: một chương trình, bảng A (hạng voucher) + bảng B (kiểm tra hiệu quả)
  · Activation-Merch: chương trình branding, không đo ROI doanh thu

NỀN SỐ: "Base/Target Sales gross" GỒM VAT — cùng nền với `net` của POS (cột `Tổng tiền`),
nên target doanh thu so thẳng với thực tế. Pre-Analysis quy "Incr Net" về ex-VAT (÷1,08)
rồi × (1 − COGS%); tương đương trên nền gồm VAT: cm_pct = (1 − COGS%) / 1,08.
"""
from __future__ import annotations

import os
import re
import unicodedata

VAT = 1.08

# loại chương trình: data_contract.json → $campaign.pre_kinds (sheet · tiền tố ID · cơ chế · nhãn)
def _contract_kinds():
    import json
    here = os.path.dirname(os.path.abspath(__file__))
    with open(os.path.join(here, "..", "data_contract.json"), encoding="utf-8") as f:
        return json.load(f)["$campaign"]["pre_kinds"]


_KINDS = _contract_kinds()
KIND_LABEL = {k["code"]: k["label"] for k in _KINDS}
KIND_MECH = {k["code"]: k["mechanic"] for k in _KINDS}
# 4 sheet có bảng chuẩn (dòng tiêu đề ID + Chương trình); Voucher · Activation đọc riêng
KIND_SHEETS = {k["sheet"]: (k["prefix"], k["code"], k["mechanic"]) for k in _KINDS
               if k["code"] not in ("VOUCHER", "ACTIVATION")}

# tiêu đề cột trong file (đã bỏ dấu, gộp khoảng trắng) → khoá chuẩn
_HDR = [
    ("id", r"^id$"), ("name", r"^chuong trinh"), ("brand", r"^brand"), ("plan_status", r"^status"),
    ("est_tc", r"^est\.? ?tc"), ("base_gross", r"^base sales"), ("growth", r"^growth"),
    ("target_gross", r"^target sales"), ("net_exvat", r"^net ex"), ("cogs_pct", r"^cogs"),
    ("gross_margin", r"^gross margin"), ("incr_net_exvat", r"^incr net"), ("incr_gm", r"^incr gm"),
    ("driver", r"^driver"), ("promo_cost", r"^promo"), ("fixed_cost", r"^fixed"),
    ("total_cost", r"^total cost"), ("net_contrib", r"^net contrib"), ("roi", r"^roi"),
    ("promo_pct_net", r"^ct%"), ("breakeven_incr", r"^breakeven"), ("assessment", r"^danh gia"),
]


def _plain(s):
    s = unicodedata.normalize("NFD", str(s or "")).replace("đ", "d").replace("Đ", "D")
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


def _num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _rows(ws):
    """Các dòng của sheet, bỏ ô trống ở mép trái TỪNG dòng: tiêu đề sheet nằm cột A còn
    bảng số đặt từ cột B — so theo vị trí tuyệt đối là lệch một cột."""
    out = []
    for r in ws.iter_rows(values_only=True):
        r = list(r)
        while r and r[0] is None:
            r.pop(0)
        out.append(r)
    return out


def _kind_table(rows):
    """Bảng chuẩn của sheet loại: tìm dòng tiêu đề có ID + Chương trình."""
    out, hdr = [], None
    for r in rows:
        cells = [_plain(c) for c in r]
        if hdr is None:
            if "id" in cells and any(c.startswith("chuong trinh") for c in cells):
                hdr = {}
                for j, c in enumerate(cells):
                    for key, rx in _HDR:
                        if key not in hdr and c and re.search(rx, c):
                            hdr[key] = j
                            break
            continue
        j = hdr.get("id")
        rid = r[j] if j is not None and j < len(r) else None
        if not rid or not re.match(r"^[A-Z]\d+$", str(rid).strip()):
            if out:
                break          # hết bảng (dòng ghi chú bên dưới)
            continue
        rec = {k: (r[i] if i < len(r) else None) for k, i in hdr.items()}
        for k in list(rec):
            if k not in ("id", "name", "brand", "plan_status", "assessment"):
                rec[k] = _num(rec[k])
        out.append(rec)
    return out


def _voucher(rows):
    """Voucher-Loyalty: tổng chi phí redemption (bảng A) + kiểm tra hiệu quả (bảng B)."""
    txt = {_plain(r[0]): r for r in rows if r and r[0] is not None}
    name = None
    for r in rows[:3]:
        m = re.search(r"\n?(.+?)\s*\.\s*Logic", str(r[0] if r else ""))
        if m:
            name = m.group(1).strip()
    total = next((r for k, r in txt.items() if k == "tong"), None)
    nums = [x for x in (total or []) if isinstance(x, (int, float))]

    def val(prefix):
        r = next((r for k, r in txt.items() if k.startswith(prefix)), None)
        return _num(next((x for x in (r or [])[1:] if isinstance(x, (int, float))), None))

    cost = val("cp redemption") or (nums[0] if nums else None)
    return dict(name=name or "Voucher / Loyalty", promo_cost=cost, fixed_cost=0.0,
                total_cost=cost, rev_linked_gross=(nums[1] if len(nums) > 1 else None),
                incr_share=val("% doanh thu that su"), cogs_pct=val("cogs % binh quan"),
                incr_net_exvat=val("dt incremental"), incr_gm=val("bien gop tang them"),
                net_contrib=val("dong gop rong"), roi=val("roi"))


def _activation(rows):
    out, hdr = [], None
    for r in rows:
        cells = [_plain(c) for c in r]
        if hdr is None:
            if cells and cells[0].startswith("chuong trinh") and any("tong cp" in c for c in cells):
                hdr = cells
            continue
        if not r or r[0] is None:
            if out:
                break
            continue
        rec = dict(zip(hdr, r))
        if str(r[0]).lower().startswith("ghi chu"):
            break
        out.append(dict(name=str(r[0]).strip(), brand=rec.get("brand"),
                        promo_cost=_num(rec.get("giai thuong/qua (vnd)")),
                        fixed_cost=_num(rec.get("cp san xuat (vnd)")),
                        total_cost=_num(rec.get("tong cp")),
                        reach_target=_num(rec.get("reach/engage muc tieu")),
                        net_contrib=-(_num(rec.get("tong cp")) or 0), roi=None))
    return out


def read(path):
    """→ list[dict] mỗi chương trình một dòng, khoá `pre_id` (C1, D2, G7, V1, A1…)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        plan = []
        master = {}
        ms = next((n for n in wb.sheetnames if "tong hop" in _plain(n)), None)
        if ms:
            for r in _rows(wb[ms]):
                if r and isinstance(r[0], (int, float)) and len(r) > 1 and r[1]:
                    master[_plain(r[1])] = dict(brand=r[2], kind_label=r[3], roi_master=r[5])
        for sheet, (prefix, kind, mech) in KIND_SHEETS.items():
            if sheet not in wb.sheetnames:
                continue
            for rec in _kind_table(_rows(wb[sheet])):
                rec.update(pre_id=str(rec.pop("id")).strip(), kind=kind, mechanic_hint=mech, sheet=sheet)
                plan.append(rec)
        vs = next((n for n in wb.sheetnames if _plain(n).startswith("voucher")), None)
        if vs:
            v = _voucher(_rows(wb[vs]))
            m = next((x for k, x in master.items() if _plain(v["name"]).split(" (")[0] in k), None)
            brand = m["brand"] if m else None
            if not m:
                m = next(((k, x) for k, x in master.items() if "voucher" in _plain(x["kind_label"])), None)
                if m:
                    v["name"] = next(r[1] for r in _rows(wb[ms]) if r and len(r) > 1 and _plain(r[1]) == m[0])
                    brand = m[1]["brand"]
            v.update(pre_id="V1", kind="VOUCHER", mechanic_hint=KIND_MECH["VOUCHER"], sheet=vs, brand=brand)
            plan.append(v)
        acs = next((n for n in wb.sheetnames if _plain(n).startswith("activation")), None)
        if acs:
            for i, a in enumerate(_activation(_rows(wb[acs])), 1):
                a.update(pre_id=f"A{i}", kind="ACTIVATION", mechanic_hint=KIND_MECH["ACTIVATION"], sheet=acs)
                # tên ngắn khớp sheet tổng hợp
                full = next((r[1] for r in _rows(wb[ms]) if ms and r and len(r) > 1 and r[1]
                             and _plain(r[1]).split(" (")[0] == _plain(a["name"]).split(" (")[0]), None)
                if full:
                    a["name"] = full
                plan.append(a)
    finally:
        wb.close()

    for p in plan:
        p["name"] = str(p.get("name") or "").strip()
        p["brand"] = str(p.get("brand") or "").strip().upper() or None
        p["kind_label"] = KIND_LABEL[p["kind"]]
        base, tgt = p.get("base_gross"), p.get("target_gross")
        # doanh thu tăng thêm trên nền GỒM VAT — cùng nền với POS
        p["incr_gross"] = (tgt - base) if (base is not None and tgt is not None) else None
        cogs = p.get("cogs_pct")
        p["cm_pct"] = ((1 - cogs) / VAT) if cogs is not None else None
        if p.get("total_cost") is None and p.get("promo_cost") is not None:
            p["total_cost"] = (p.get("promo_cost") or 0) + (p.get("fixed_cost") or 0)
        if p.get("est_tc") and tgt:
            p["target_aov"] = tgt / p["est_tc"]
        p["source_file"] = os.path.basename(path)
    return plan


# ═════════════ FILE DECK THEO QUÝ (từ Q4/2026) → MẪU CHUẨN M7.1 ═════════════
# File kế hoạch quý dựng từ deck marketing — 4 sheet, tên có tiền tố số + mã quý:
#   `03_Q4 Master Plan`  mỗi chương trình một dòng (brand · loại · timeline · outlet · cơ chế · KPI · chi phí MKT)
#   `04_Q4 Pre-Analysis` business case của deck (BRAND · PROGRAM · METRIC × 3 kịch bản)
#   `05_Q4 Budget` · `06_Q4 Calendar`
# Hệ thống KHÔNG dựng màn hình riêng cho định dạng này: `deck_programs()` chuyển từng chương trình sang
# ĐÚNG các cột sheet `chuong_trinh` của sổ chuẩn ($preeval.input_fields), rồi tools/preeval.py đánh giá
# như mọi chương trình khác trên M7.1. Chỉ điền ô deck ghi RÕ (ngày dd/mm, tên cửa hàng nhận ra được);
# ô không rõ để TRỐNG, kèm gợi ý "deck ghi gì" để người dùng bổ sung ở sổ (cùng program_id).

_QHDR = [
    ("brand", r"^brand$"), ("name", r"^program"), ("type", r"^type$"), ("timeline", r"^timeline"),
    ("outlet", r"^outlet"), ("objective", r"^objective"), ("mechanic", r"^mechanic"),
    ("rationale", r"^source analysis"), ("kpi", r"^primary kpi"), ("mkt_cost", r"^mkt cost"), ("note", r"^note$"),
]
_BRANDING_RX = re.compile(r"brand|key window|guest shift|performance")


def _sheet_like(wb, key):
    """Sheet có tên chứa `key` (bỏ tiền tố số + mã quý: `03_Q4 Master Plan` → master plan)."""
    return next((n for n in wb.sheetnames if key in _plain(n)), None)


def file_format(path):
    """'QUARTER' (file deck theo quý, từ Q4/2026) · 'LEGACY' (6 sheet loại của Q3) · None."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if _sheet_like(wb, "master plan"):
            return "QUARTER"
        if any("tong hop" in _plain(n) for n in wb.sheetnames):
            return "LEGACY"
        return None
    finally:
        wb.close()


def files_by_format():
    """Mọi file S16 xếp theo định dạng → {'LEGACY': […], 'QUARTER': […], None: […], 'ERROR': […]}.
    Nơi DUY NHẤT chọn file kế hoạch: chọn theo định dạng, không theo ngày sửa — file quý mới
    nhất không được đè mất kế hoạch Q3 đang nối M7.2 (campaign.py) và pre_q3 (build_mkt.py)."""
    from monthly_lib import l0_files
    out = {"LEGACY": [], "QUARTER": [], None: [], "ERROR": []}
    for f in l0_files("S16_pre_analytics"):
        try:
            out[file_format(f)].append(f)
        except Exception:  # noqa: BLE001 — file hỏng / đang chép dở
            out["ERROR"].append(f)
    return out


def latest_legacy():
    """File định dạng 6 sheet loại (Q3/2026) sửa gần nhất."""
    fs = files_by_format()["LEGACY"]
    return max(fs, key=os.path.getmtime) if fs else None


def _table(ws, spec, must):
    """Bảng có dòng tiêu đề khớp `spec` → list[dict]."""
    out, hdr = [], None
    for r in ws.iter_rows(values_only=True):
        r = list(r)
        cells = [_plain(c) for c in r]
        if hdr is None:
            h = {}
            for j, c in enumerate(cells):
                for key, rx in spec:
                    if key not in h and c and re.search(rx, c):
                        h[key] = j
                        break
            if all(k in h for k in must):
                hdr = h
            continue
        if any(v is not None for v in r[:max(hdr.values()) + 1]):
            out.append({k: (r[i] if i < len(r) else None) for k, i in hdr.items()})
    return out


def _long_rows(ws):
    """Bảng dạng dài có tiêu đề ở dòng bắt đầu bằng BRAND → (tiêu đề đã bỏ dấu, các dòng)."""
    rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows) if r and _plain(r[0]) == "brand"), None)
    if hi is None:
        return [], []
    return [_plain(c) for c in rows[hi]], [r for r in rows[hi + 1:] if r and r[0]]


def _compact(s):
    return re.sub(r"[^a-z0-9]", "", _plain(s))


def _same_program(a, b):
    """Tên ngắn ở sheet Pre-Analysis/Calendar ↔ tên đầy đủ ở Master Plan
    (`Her Wonders 20/10` ↔ `Her Wonders – Welcome Mocktail 20/10`). → điểm 0–1."""
    ca, cb = _compact(a), _compact(b)
    if not ca or not cb:
        return 0.0
    if ca == cb:
        return 1.0
    if ca in cb or cb in ca:
        return 0.95
    ta = {t for t in re.split(r"[^a-z0-9]+", _plain(a)) if t}
    tb = {t for t in re.split(r"[^a-z0-9]+", _plain(b)) if t}
    return len(ta & tb) / min(len(ta), len(tb)) if ta and tb else 0.0


def _best(name, brand, cands, floor=0.6):
    best, score = None, floor
    for c in cands:
        if c["brand"] not in (brand, "ALL"):
            continue
        s = _same_program(name, c["name"])
        if s > score or (s == score and best is None):
            best, score = c, s
    return best


def _stores(outlet, brand):
    """Chuỗi OUTLET của deck → (mã cửa hàng 'A|B' hoặc None, gợi ý khi không nhận ra).
    Nhận ra: tên có alias ở dim_store (39NTMK · The Berkley · SSV · The Crest · SKC…) và cụm cả brand
    (`2 outlets` · `NDC outlets` · `NDC` · `All NOIRE system` → các cửa hàng flagship/core của brand).
    Chỉ nhận ra một phần (vd `ET + TM + SKC`) → để trống, không điền thiếu cửa hàng."""
    from monthly_lib import STORE_META, store_in_text
    raw = str(outlet or "").strip()
    t = _plain(raw)
    if not t:
        return None, "deck không ghi cửa hàng"
    if "khong chi ro" in t:
        return None, f"deck ghi '{raw}'"
    own = [c for c, m in STORE_META.items() if m.get("brand") == brand and m.get("tier") in ("flagship", "core")]
    head = re.sub(r"\(.*?\)", "", t).strip()
    if re.fullmatch(rf"(\d+ outlets?|{brand.lower()}( outlets?)?|all( noire)?( system)?)", head):
        n = re.match(r"(\d+) outlet", head)
        if own and (not n or int(n.group(1)) == len(own)):
            return "|".join(own), None
        return None, f"deck ghi '{raw}' — brand có {len(own)} cửa hàng"
    parts = [x.strip() for x in re.split(r"\+|–|—|,|;|/|\s-\s", raw) if x.strip()]
    codes, bad = [], []
    for x in parts:
        c = store_in_text(x)
        if c:
            codes.append(c)
        elif x != x.lower():      # cụm mô tả viết thường (`shared decor concept`) không phải tên cửa hàng
            bad.append(x)
    if codes and not bad:
        return "|".join(dict.fromkeys(codes)), None
    return None, f"deck ghi '{raw}'" + (f" — chưa nhận ra: {', '.join(bad)}" if bad else "")


_RANGE = re.compile(r"(\d{1,2})/(\d{1,2})\s*[–—-]\s*(\d{1,2})/(\d{1,2})")
_RANGE_M = re.compile(r"(?<![\d/])(\d{1,2})\s*[–—-]\s*(\d{1,2})/(\d{1,2})(?![\d/])")
_FULL = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})")


def _dates(timeline, year):
    """Timeline deck → (date_from, date_to, gợi ý). Chỉ nhận ngày ghi RÕ dd/mm; `Tháng 10`, `Q4`,
    `Christmas Q4` → để trống."""
    from datetime import date
    raw = str(timeline or "").strip()
    if not raw:
        return None, None, "deck không ghi thời gian"
    if not year:
        return None, None, f"deck ghi '{raw}' — không rõ năm (tên file thiếu Q#_YYYY)"
    try:
        m = _RANGE.search(raw)
        if m:
            d0 = date(year, int(m.group(2)), int(m.group(1)))
            d1 = date(year, int(m.group(4)), int(m.group(3)))
        else:
            m = _RANGE_M.search(raw)
            if m:
                d0 = date(year, int(m.group(3)), int(m.group(1)))
                d1 = date(year, int(m.group(3)), int(m.group(2)))
            else:
                m = _FULL.search(raw)
                if not m:
                    return None, None, f"deck ghi '{raw}' — chưa có ngày cụ thể"
                d0 = d1 = date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None, None, f"deck ghi '{raw}' — ngày không hợp lệ"
    if d1 < d0:
        d1 = d1.replace(year=d1.year + 1)
    many = len(re.findall(r"\d{1,2}/\d{1,2}", raw)) > (2 if _RANGE.search(raw) else 1)
    return d0.isoformat(), d1.isoformat(), (f"deck ghi '{raw}' — có nhiều ngày, kiểm lại" if many else None)


def _dows(timeline):
    """`T2–T6` · `T2,T3,T5,CN` → '0|1|2|3|4' (0 = thứ 2 … 6 = CN). Không ghi thứ → None."""
    t = str(timeline or "")
    out = set()
    for a, b in re.findall(r"\bT([2-7])\s*[–—-]\s*T([2-7])\b", t):
        out.update(range(int(a) - 2, int(b) - 1))
    for x in re.findall(r"\b(T[2-7]|CN)\b", t):
        out.add(6 if x == "CN" else int(x[1]) - 2)
    return "|".join(str(x) for x in sorted(out)) or None


def _slug(name):
    head = re.split(r"\s[–—/-]\s|\(", str(name))[0]
    return re.sub(r"[^A-Z0-9]", "", _plain(head).upper())[:14] or "CT"


def _vn(x, unit):
    """Số deck → chữ đọc được cho ghi chú: 135.683.034 VND → '135,7 tr' · 0,45 % → '45%'."""
    u = _plain(unit)
    if x is None:
        return "—"
    if u == "%":
        return f"{x * 100:.0f}%" if abs(x) <= 1.5 else f"{x:.0f}%"
    if u == "vnd" or abs(x) >= 1e6:
        return f"{x / 1e6:.1f} tr".replace(".", ",")
    s = f"{x:.2f}".rstrip("0").rstrip(".").replace(".", ",")
    return s + ("×" if u == "x" else f" {unit}" if unit else "")


def deck_to_standard(path):
    """Một file deck quý → dict(quarter, programs, costs, issues). `programs` mang đúng cột sheet
    `chuong_trinh` của sổ chuẩn + `_hint` (deck ghi gì ở ô còn trống) + `_quarter` · `_file`."""
    from openpyxl import load_workbook
    b = os.path.basename(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ms = _sheet_like(wb, "master plan")
        title = str(next(wb[ms].iter_rows(values_only=True), [None])[0] or "")
        m = re.search(r"Q([1-4])[ _.-]*(\d{4})", b, re.I) or re.search(r"Q([1-4])[ _.-]*(\d{4})", title, re.I)
        qn, year = (int(m.group(1)), int(m.group(2))) if m else (None, None)
        quarter = f"{year}-Q{qn}" if qn else None
        master = _table(wb[ms], _QHDR, must=("brand", "name"))

        cases = {}                                   # business case deck (Cơ sở) → ghi chú tham khảo
        pa = _sheet_like(wb, "pre-analysis")
        if pa:
            h, rows = _long_rows(wb[pa])
            base = next((j for j, c in enumerate(h) if c == "co so"), None)
            unit = h.index("unit") if "unit" in h else None
            for r in rows:
                if len(r) > 2 and r[2] and base is not None and base < len(r):
                    u = r[unit] if unit is not None and unit < len(r) else None
                    cases.setdefault((str(r[0]).strip().upper(), str(r[1]).strip()), []).append(
                        f"{r[2]} {_vn(_num(r[base]), u)}")
        cal = {}
        cs = _sheet_like(wb, "calendar")
        if cs:
            h, rows = _long_rows(wb[cs])
            mcols = [(j, c) for j, c in enumerate(h) if c[:3] in ("oct", "nov", "dec", "jan", "feb", "mar",
                                                                   "apr", "may", "jun", "jul", "aug", "sep")]
            for r in rows:
                marks = [f"{c[:3].title()}{'' if str(r[j]).strip() == '●' else ' ' + str(r[j]).strip()}"
                         for j, c in mcols if j < len(r) and r[j]]
                cal[(str(r[0]).strip().upper(), str(r[1]).strip())] = ", ".join(marks)
    finally:
        wb.close()

    programs, costs, issues, seen, used = [], [], [], set(), set()
    for i, row in enumerate(master, 1):
        brand = str(row.get("brand") or "").strip().upper()
        name = str(row.get("name") or "").strip()
        if not brand or not name:
            continue
        pid = f"{brand}-{year}Q{qn}-{_slug(name)}" if qn else f"{brand}-{_slug(name)}"
        while pid in seen:
            pid += "2"
        seen.add(pid)
        hint = {}
        stores, hint["store_scope"] = _stores(row.get("outlet"), brand)
        d0, d1, dn = _dates(row.get("timeline"), year)
        hint["date_from"] = hint["date_to"] = dn
        branding = bool(_BRANDING_RX.search(_plain(row.get("type"))))
        hint["benefit"] = f"deck ghi: {row.get('mechanic')}" if row.get("mechanic") else "deck không ghi cơ chế"
        hint["condition"] = hint["benefit"]
        hint["objective"] = f"deck: {row.get('objective')}" if row.get("objective") and not branding else None
        hint["lever_primary"] = hint["objective"]
        ck = _best(name, brand, [dict(brand=k[0], name=k[1], key=k) for k in cases if k not in used])
        if ck:
            used.add(ck["key"])
            hint["est_bills"] = "deck (Cơ sở): " + " · ".join(cases[ck["key"]][:3])
        ca = _best(name, brand, [dict(brand=k[0], name=k[1], key=k) for k in cal])
        note = [f"Nhập tự động từ {b} (Master Plan dòng {i})", f"Loại: {row.get('type')}"]
        if ca and cal[ca["key"]]:
            note.append(f"Calendar: {cal[ca['key']]}")
        if row.get("rationale"):
            note.append(f"Cơ sở phân tích: {row.get('rationale')}")
        if row.get("kpi"):
            note.append(f"KPI: {row.get('kpi')}")
        if ck:
            note.append("Business case deck (Cơ sở): " + " · ".join(cases[ck["key"]]))
        if row.get("note"):
            note.append(f"Ghi chú deck: {row.get('note')}")
        programs.append(dict(
            program_id=pid, name=name, brand=brand, store_scope=stores, date_from=d0, date_to=d1,
            dow=_dows(row.get("timeline")), objective="BRANDING" if branding else None, lever_primary=None,
            content=row.get("mechanic"), hypothesis=row.get("objective"), status="NHAP",
            note=" · ".join(str(x) for x in note), _quarter=quarter, _file=b,
            _hint={k: v for k, v in hint.items() if v}))
        mc = _num(row.get("mkt_cost"))
        if mc:
            costs.append(dict(program_id=pid, cost_type="OTHER", amount=mc, vat_pct=None,
                              note=f"Chi phí MKT theo deck ({b})"))
    for k in cases:
        if k not in used:
            issues.append(f"{b}: business case '{k[1]}' ({k[0]}) không khớp chương trình nào ở Master Plan")
    return dict(quarter=quarter, programs=programs, costs=costs, issues=issues)


def deck_programs():
    """Mọi file deck quý (mỗi quý lấy file sửa gần nhất) → dict(programs, costs, issues) theo mẫu chuẩn."""
    by_q, out = {}, dict(programs=[], costs=[], issues=[])
    fm = files_by_format()
    for f in fm["ERROR"]:
        out["issues"].append(f"{os.path.basename(f)}: không mở được (hỏng hoặc đang chép dở)")
    for f in fm["QUARTER"]:
        try:
            d = deck_to_standard(f)
        except Exception as e:  # noqa: BLE001 — một file hỏng không chặn cả M7.1
            out["issues"].append(f"{os.path.basename(f)}: không đọc được — {str(e)[:80]}")
            continue
        k = d["quarter"] or os.path.basename(f)
        if k not in by_q or os.path.getmtime(f) > by_q[k][0]:
            by_q[k] = (os.path.getmtime(f), d)
    for _, d in sorted(by_q.values(), key=lambda x: str(x[1]["quarter"])):
        for k in out:
            out[k] += d[k]
    return out


if __name__ == "__main__":
    import sys
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    for _s in (sys.stdout, sys.stderr):
        _s.reconfigure(encoding="utf-8", errors="replace")
    f = latest_legacy()
    if f:
        print("══", os.path.basename(f), "(kế hoạch Q3 → M7.2)")
        for p in read(f):
            print(" ", p["pre_id"], p["brand"], p["kind"], p["name"][:40], p.get("target_gross"), p.get("roi"))
    d = deck_programs()
    for p in d["programs"]:
        print(p["program_id"], "|", p["store_scope"], p["date_from"], p["date_to"], p["dow"], p["objective"], "|",
              {k: v[:70] for k, v in p["_hint"].items() if k in ("store_scope", "date_from")})
    for x in d["issues"]:
        print("  !", x)
