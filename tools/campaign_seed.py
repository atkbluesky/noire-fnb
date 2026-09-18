# -*- coding: utf-8 -*-
"""
DỰNG FILE KHAI BÁO THẬT M7 PROMOTION — Campaign_Tracking_2026.xlsx
==================================================================
    python tools/campaign_seed.py            tạo file (từ chối nếu file đã có)
    python tools/campaign_seed.py --merge    file đã có: CHỈ THÊM dòng mới (tên CTKM mới trên POS,
                                             chương trình kế hoạch mới) — không đụng ô đã nhập
    python tools/campaign_seed.py --dry      in ra sẽ ghi gì, không ghi

Gom MỌI tên chương trình của cụm Promotion vào MỘT danh mục — master chung cho
M7 Promotion · M7.1 Pre-Analytics · M7.2 Promotion Tracking:

  ① Kế hoạch   L0_input/03_MARKETING/05_Promotion_Ke_Hoach   (S16 · Pre-Analysis — đọc qua tools/pre_analysis.py)
  ② Thực tế    L0_input/01_DOANH_THU/03_POS_Hoa_Don          (S02 · cột `Tên CTKM` trên hoá đơn → fact_promo_day)
  ③ Báo cáo    L0_input/03_MARKETING/06_Promotion_Ket_Qua    (S17 · báo cáo hiệu quả LTO)

NGUYÊN TẮC: MỖI CHƯƠNG TRÌNH MỘT TÊN RIÊNG → MỖI TÊN CTKM TRÊN POS LÀ MỘT DÒNG. Không gom các tên
gần giống nhau, không đoán ghép. Kế hoạch ↔ POS chỉ ghép ở bảng PRE_MATCH (team xác nhận từng dòng).
Một kế hoạch có nhiều mức trên POS (Noire Passport 50K/100K/150K) → mỗi mức một dòng, cùng pre_id,
engine chấm chung một lần.

Màu ô trong file:
  CAM  = cần bạn nhập (nội dung, giả thuyết, mục tiêu · phương án, người phụ trách, ai trả tiền…)
  XANH = máy điền từ dữ liệu, cần kiểm lại (ngày/cửa hàng theo lần đầu-cuối thấy trên POS,
         mục tiêu · phương án · cơ chế đoán theo tên)
  trắng = số liệu gốc (tên POS, tên kế hoạch, mã kế hoạch)
"""
from __future__ import annotations

import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from monthly_lib import (  # noqa: E402
    BRAND_OF_STORE, CONTRACT, SOURCES, classify_nature, l0_dir, l0_latest, store_code,
)
import campaign as E  # noqa: E402
import campaign_template as TPL  # noqa: E402
import pre_analysis as PA  # noqa: E402

C = CONTRACT["$campaign"]
OUT_NAME = "Campaign_Tracking_2026.xlsx"
SKIP_NATURE = {"INTERNAL"}          # ưu đãi nội bộ không phải chương trình marketing

# ─────────────── ghép kế hoạch (pre_id) ↔ tên POS — ĐÃ XÁC NHẬN ───────────────
#   pre_id → danh sách regex, mỗi regex khớp ĐÚNG một tên CTKM trên POS
#   D3 Pre-booking 15% KHÔNG phải BERKLEY - PRE-BOOKING OFFER (team xác nhận 17/09/2026) — không ghép.
PRE_MATCH = {
    "G1": ([r"^dining - the monday treat$"], "NDC · tặng dessert thứ 2 · POS NDC_NTMK từ 17/08"),
    "G3": ([r"^tặng quà sinh nhật \(bánh millie crepe\)$"], "NDC · tặng bánh sinh nhật · POS NDC_NTMK từ 02/08"),
    "G7": ([r"^complimemtary tataki wagyu$"], "NJFB · tặng Tataki Wagyu nhóm ≥4 · POS 02–29/08"),
    "G8": ([r"^jfb - obon table$", r"^obon table$"], "NJFB · Obon Table · POS NJFB_SSV 26–30/08"),
    "V1": ([r"^noire passport \(50k\)$", r"^noire passport \(100k\)$", r"^noire passport \(150k\)$"],
           "NCB · voucher bậc thang — mỗi mức một dòng, chấm chung kế hoạch V1"),
}

DOW_LABEL = ["Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6", "Thứ 7", "Chủ nhật"]


def log(*a):
    print(*a, flush=True)


def _low(s):
    return str(s or "").strip().lower()


def _slug(title, n=14):
    s = unicodedata.normalize("NFD", title).replace("đ", "d").replace("Đ", "D")
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn").upper()
    words = [w for w in re.split(r"[^A-Z0-9]+", s) if w and w not in ("THE", "NOIRE", "DC", "JFB", "CHO", "VA")]
    out = ""
    for w in words:
        if len(out) + len(w) > n and out:
            break
        out += w
    return out[:n] or "CT"


# ─────────────── đoán cơ chế / cửa sổ / nhịp (ô XANH) ───────────────
MECH_RULES = [
    ("LUCKY_DRAW", r"lucky draw"),
    ("BXGY", r"1 tặng 1|buy 01 .*get|mua 1 tặng 1"),
    ("COMBO_SET", r"đồng giá|combo|set menu|bánh cake \d+k|bánh pastry \d+k|power lunch"),
    ("CASH_VOUCHER", r"voucher|passport"),
    ("MERCH_GIFT", r"tặng khăn|mug|muse gift"),
    ("GIFT_ITEM", r"tặng|complim|gift|table$|sweet|love bloom|refined hour|morning delight"),
    ("PCT_OFF", r"\d+ ?%|% off|off the"),
    ("FIXED_OFF", r"\d+k\b"),
]
# (objective, lever) — slide 3 (6 phương án TC) · slide 4 (3 phương án AOV) · Branding
LEVER_RULES = [
    ("TC", "TC_BASE", r"lucky draw|đăng ký mới|soft opening|pre-booking|đặt bàn trước"),   # trúng thưởng = Customer Base (PDF slide 3)
    ("BRANDING", "BR_REVIEW", r"google review|khảo sát ý kiến|cskh"),
    ("BRANDING", "BR_ENGAGE", r"check-in|social media|minigame|mini game"),
    ("TC", "TC_FREQ", r"passport|loyalty|hạng |thẻ vip|voucher kỳ sau|members"),
    ("TC", "TC_PARTY", r"nhóm|party|table$"),
    ("TC", "TC_DAYPART", r"happy hour|lunch|breakfast|brunch|morning|refined hour|thứ \d|tuesday|wednesday|thursday|monday"),
    ("TC", "TC_CHANNEL", r"grab|dine out|delivery|online"),
    ("AOV", "AOV_QTY", r"combo|set menu|đồng giá|1 tặng 1|buy 01|mua 1 tặng 1"),
    ("AOV", "AOV_VALUE", r"wagyu|upsell|cao cấp"),
]


def guess_lever(text):
    t = _low(text)
    return next(((o, lv) for o, lv, rx in LEVER_RULES if re.search(rx, t)), (None, None))


WINDOW_RULES = [
    ("TET_AL", r"tết"), ("INT_WOMEN_0803", r"08/03|8/3"), ("VALENTINE", r"valentine|love bloom"),
    ("NATIONAL_0209", r"02/09|2/9|firework|độc lập|independence"), ("VN_WOMEN_2010", r"20/10"),
    ("SUMMER", r"summer|hè"), ("OPENING", r"soft opening|berkley"),
]


def guess(rx_list, text):
    t = _low(text)
    return next((code for code, rx in rx_list if re.search(rx, t)), None)


# ─────────────── ② POS ───────────────
def pos_groups(promo, last):
    names = defaultdict(lambda: dict(stores=set(), dates=set(), nature=Counter(), bills=0.0, net=0.0))
    for r in promo:
        nm = str(r.get("name_pos") or "").strip()
        if not nm:
            continue
        g = names[nm]
        g["stores"].add(r["store"])
        g["dates"].add(r["date"])
        g["nature"][r.get("nature") or classify_nature(nm)] += 1
        g["bills"] += E.to_num(r.get("bills"), 0) or 0
        g["net"] += E.to_num(r.get("net"), 0) or 0
    groups = {}
    for nm, g in names.items():
        nat = g["nature"].most_common(1)[0][0]
        if nat in SKIP_NATURE:
            continue
        key, title = "N:" + nm, nm                     # mỗi tên CTKM = một chương trình
        x = groups.setdefault(key, dict(key=key, title=title, names=[], stores=set(),
                                        dates=set(), nature=Counter(), bills=0.0, net=0.0))
        x["names"].append(nm)
        x["stores"] |= g["stores"]
        x["dates"] |= g["dates"]
        x["nature"][nat] += len(g["dates"])
        x["bills"] += g["bills"]
        x["net"] += g["net"]
    for x in groups.values():
        x["names"].sort()
        # brand theo cửa hàng của 3 thương hiệu; cửa hàng ngoài (IFC_SIG = OTHER) không làm thành 'ALL'
        x["brands"] = sorted({BRAND_OF_STORE.get(s) for s in x["stores"]} & {"NCB", "NDC", "NJFB"}) or ["ALL"]
        x["first"], x["last"] = min(x["dates"]), max(x["dates"])
        # còn thấy trên POS trong 7 ngày cuối (chương trình theo tuần mới lặp 1 lần/tuần) → đang chạy
        x["running"] = (last - x["last"]).days <= 7
    return groups


# ─────────────── ③ báo cáo hiệu quả LTO (S17) ───────────────
def report_lto():
    f = l0_latest("S17_lto_actual")
    if not f:
        return None
    from openpyxl import load_workbook
    wb = load_workbook(f, read_only=True, data_only=True)
    try:
        def rows(prefix):
            n = next((s for s in wb.sheetnames if PA._plain(s).startswith(prefix)), None)
            return PA._rows(wb[n]) if n else []

        plan = {PA._plain(r[0]): r for r in rows("4.") if r and isinstance(r[0], str)}
        m = re.search(r"(\d{4})", os.path.basename(f))
        year = int(m.group(1)) if m else date.today().year
        span = re.findall(r"(\d{2})/(\d{2})", str((plan.get("thoi gian chay") or [None, None, ""])[2]))
        d0 = date(year, int(span[0][1]), int(span[0][0])) if len(span) >= 2 else None
        d1 = date(year, int(span[1][1]), int(span[1][0])) if len(span) >= 2 else None
        stores = []
        for r in rows("5."):
            # dòng cửa hàng: cột 1 = tên cửa hàng, cột 2 = brand (bỏ dòng TỔNG, dòng tiêu đề)
            if len(r) > 1 and isinstance(r[0], str) and str(r[1]).strip() in ("NCB", "NDC", "NJFB"):
                code = store_code(r[0])
                if code and code not in stores:
                    stores.append(code)

        def pa(label):
            r = plan.get(label) or []
            return (PA._num(r[1]) if len(r) > 1 else None), (PA._num(r[2]) if len(r) > 2 else None)

        gift_p, gift_a = pa("chi phi qua")
        media_p, media_a = pa("chi phi truyen thong")
        cups_p, cups_a = pa("so ly lto ban ra")
        rev_p, _ = pa("doanh thu (chi ly lto)")
        items, on = [], False           # chỉ bảng "Món LTO" đầu sheet 2, dừng ở dòng TỔNG
        for r in rows("2."):
            if r and isinstance(r[0], str) and PA._plain(r[0]).startswith("mon lto"):
                on = True
                continue
            if on:
                if not r or not isinstance(r[0], str) or str(r[0]).upper().startswith("TỔNG"):
                    break
                items.append(str(r[0]).strip())
        return dict(file=os.path.basename(f), date_from=d0, date_to=d1, stores=stores,
                    gift=(gift_p, gift_a), media=(media_p, media_a), cups=(cups_p, cups_a), rev_plan=rev_p,
                    items=list(dict.fromkeys(items)))
    finally:
        wb.close()


# ─────────────── dựng dòng ───────────────
DIM_KEYS = [c for c, _ in TPL.DIM_COLS]
NEED_DIM = ("content", "hypothesis", "objective", "lever_primary", "owner", "cost_owner", "discount_rule",
            "mechanic", "store_scope", "date_from")


def build(plan, groups, rep, last):
    dims, tgts, costs, fills = [], [], [], {}
    used_ids = set()

    def cid_for(brand, d0, title):
        base = f"{brand}-{d0:%Y-%m}" if d0 else f"{brand}-2026-Q3"
        cid = f"{base}-{_slug(title)}"
        k = 2
        while cid in used_ids:
            cid = f"{base}-{_slug(title, 12)}{k}"
            k += 1
        used_ids.add(cid)
        return cid

    def add(row, guessed=(), target=None, cost_lines=()):
        i = len(dims)
        for k in DIM_KEYS:
            row.setdefault(k, None)
        for k in guessed:
            if row.get(k) not in (None, ""):
                fills[("dim_campaign", i, k)] = "guess"
        if (row.get("status") or "") != "CANCELLED":
            for k in NEED_DIM:
                if row.get(k) in (None, ""):
                    fills[("dim_campaign", i, k)] = "need"
        dims.append(row)
        if target:
            j = len(tgts)
            for k in target.pop("_need", ()):
                fills[("campaign_target", j, k)] = "need"
            tgts.append(target)
        costs.extend(cost_lines)

    def pos_fields(g):
        brand = g["brands"][0] if len(g["brands"]) == 1 else "ALL"
        obj, lv = guess_lever(g["title"])
        return dict(
            name_pos=g["title"], brand=brand, store_scope="|".join(sorted(g["stores"])),
            nature=g["nature"].most_common(1)[0][0],
            date_from=str(g["first"]), date_to=None if g["running"] else str(g["last"]),
            status="RUNNING" if g["running"] else "ENDED",
            objective=obj, lever_primary=lv,
            mechanic=guess(MECH_RULES, g["title"]),
            window=guess(WINDOW_RULES, g["title"]),
        ), ["nature", "date_from", "date_to", "status", "store_scope", "objective", "lever_primary", "mechanic", "window"]

    # ① kế hoạch: dòng kế hoạch; khớp POS thì MỖI tên POS một dòng cùng pre_id
    taken = set()
    for p in plan:
        pid = p["pre_id"]
        rx, why = PRE_MATCH.get(pid, ([], ""))
        hits = [g for g in groups.values() if any(re.search(r, _low(g["title"])) for r in rx)]
        tgt_note = (f"Target + chi phí kế hoạch TỰ LẤY từ Pre-Analysis {pid} — không cần gõ lại. "
                    "Chỉ điền submitted nếu ngày nộp kế hoạch khác ngày sửa file.")
        obj, lv = guess_lever(p["name"])
        if not hits:
            row = dict(name=p["name"], pre_id=pid, brand=p.get("brand"), status="PLANNED", source="PRE",
                       nature="COMMERCIAL", mechanic=p.get("mechanic_hint"),
                       objective="BRANDING" if p["kind"] == "ACTIVATION" else obj,
                       lever_primary="BR_ENGAGE" if p["kind"] == "ACTIVATION" else lv,
                       window=guess(WINDOW_RULES, p["name"]),
                       match_note=f"Chưa thấy trên POS (kế hoạch: {p.get('plan_status') or '—'}) — khi chạy: điền tên CTKM trên POS vào name_pos, ngày, cửa hàng")
            row["campaign_id"] = cid_for(row.get("brand") or "ALL", None, p["name"])
            add(row, ["nature", "objective", "lever_primary", "window"],
                target=dict(campaign_id=row["campaign_id"], base_method="PRE_ANALYSIS", note=tgt_note))
            continue
        for g in sorted(hits, key=lambda x: x["first"]):
            taken.add(g["key"])
            pf, gs = pos_fields(g)
            row = dict(pf, name=p["name"] if len(hits) == 1 else f"{p['name']} · {g['title']}",
                       pre_id=pid, source="PRE_POS", mechanic=p.get("mechanic_hint"),
                       match_note=f"KHỚP kế hoạch {pid}: {why}")
            if obj and not pf.get("objective"):
                row.update(objective=obj, lever_primary=lv)
            if p.get("brand") and row["brand"] not in (p["brand"], "ALL"):
                row["match_note"] += f" · LƯU Ý: kế hoạch ghi {p['brand']}, POS là {row['brand']}"
            row["campaign_id"] = cid_for(row["brand"], g["first"], g["title"])
            add(row, [k for k in gs if k != "mechanic"],
                target=dict(campaign_id=row["campaign_id"], base_method="PRE_ANALYSIS", note=tgt_note))

    # ② mỗi tên CTKM trên POS chưa có kế hoạch = một dòng · ③ báo cáo LTO
    for g in sorted(groups.values(), key=lambda x: (x["first"], -x["net"])):
        if g["key"] in taken:
            continue
        pf, gs = pos_fields(g)
        row = dict(pf, name=g["title"], source="POS", match_note="Tên CTKM trên POS · chưa có trong kế hoạch Pre-Analysis")
        cost_lines, tneed = [], ("tgt_incr_net", "submitted")
        tnote = "Chưa có kế hoạch — có target nộp TRƯỚC ngày chạy thì điền; không có thì để trống (M7.2 báo CHƯA CÓ TARGET)."
        if re.search(r"^(giảm giá|chiết khấu) trực tiếp$|xử lý tình huống", _low(g["title"])):
            row.update(status="CANCELLED", objective=None, lever_primary=None,
                       match_note="Không phải chương trình (giảm tay tại quầy / xử lý khiếu nại) — để CANCELLED cho khỏi báo 'chưa khai'")
            tneed = ()
        if _low(g["title"]).startswith("summer crush") and rep:
            row.update(name="LTO Summer Crush", source="REPORT", brand="ALL",
                       store_scope="|".join(rep["stores"]) or row["store_scope"],
                       date_from=str(rep["date_from"] or row["date_from"]),
                       date_to=str(rep["date_to"] or row["date_to"]), status="ENDED", window="SUMMER",
                       mechanic="MERCH_GIFT", objective="AOV", lever_primary="AOV_VALUE",
                       content="3 món LTO (Kyoto Berry · Peachy Berry · Blush Garden) 89K. Mua ≥2 ly LTO tặng khăn Bandana; bill ≥300K có ≥1 ly LTO tặng quạt nan gỗ. Mỗi bill 1 quà.",
                       discount_rule="Mua ≥2 ly LTO tặng khăn · bill ≥300K có LTO tặng quạt",
                       match_note=f"Theo báo cáo {rep['file']} (kỳ, cửa hàng, nội dung, chi phí). "
                                  "POS chỉ có cờ 'Summer Crush (Tặng Khăn)' — doanh thu ly LTO nằm ở nhóm món LTO DRINK")
            gp, ga = rep["gift"]
            mp, ma = rep["media"]
            cost_lines = [
                dict(campaign_id=None, cost_type="MERCH", planned=gp, actual=ga,
                     note="Quà đã tặng (27 khăn + 280 quạt). Đã sản xuất đủ kế hoạch; phần tồn dùng lại được — theo báo cáo S17"),
                dict(campaign_id=None, cost_type="KOL", planned=mp, actual=0,
                     note="Kế hoạch Review TikToker — xác nhận không thực chi (báo cáo S17)"),
                dict(campaign_id=None, cost_type="ADS", planned=None, actual=ma,
                     note="11 nhóm quảng cáo Facebook gắn tên Summer Crush, T6 + T7 (báo cáo S17)"),
            ]
            cp, ca = rep["cups"]
            if cp and ca and rep.get("rev_plan"):
                tnote = (f"Kế hoạch báo cáo: {cp:,.0f} ly · doanh thu ly LTO {rep['rev_plan']:,.0f} đ tính trên GIÁ MENU "
                         f"trước phí DV + VAT — KHÁC nền doanh thu POS nên chưa điền tgt_net. Thực tế {ca:,.0f} ly.").replace(",", ".")
        row["campaign_id"] = cid_for(row["brand"], date.fromisoformat(row["date_from"]), row["name"])
        for ln in cost_lines:
            ln["campaign_id"] = row["campaign_id"]
        tgt = None
        if row["nature"] in ("COMMERCIAL", "LOYALTY", "PARTNER") and row.get("status") != "CANCELLED" \
                and row.get("objective") != "BRANDING":
            tgt = dict(campaign_id=row["campaign_id"], base_method="DOW4W", note=tnote, _need=tneed)
        add(row, gs, target=tgt, cost_lines=cost_lines)
    return dims, tgts, costs, fills


GUIDE = [
    ("M7 PROMOTION · DANH MỤC CHƯƠNG TRÌNH (master chung M7 · M7.1 · M7.2)", True),
    ("", False),
    ("DỰNG SẴN TỪ DỮ LIỆU THẬT", True),
    ("· Kế hoạch: file Pre-Analysis (03_MARKETING/05_Promotion_Ke_Hoach) — mỗi chương trình có pre_id (C1, D3, G7, V1…).", False),
    ("· Thực tế: cột 'Tên CTKM' trên file hoá đơn POS (01_DOANH_THU/03_POS_Hoa_Don) — MỖI TÊN MỘT DÒNG, đúng từng chữ.", False),
    ("· Báo cáo hiệu quả: 03_MARKETING/06_Promotion_Ket_Qua (LTO Summer Crush — kỳ, cửa hàng, nội dung, chi phí).", False),
    ("· Ưu đãi nội bộ (Chairman, Directors, nhân viên…) KHÔNG đưa vào — không phải chương trình marketing.", False),
    ("", False),
    ("CÁC CỘT NHẬP TAY — NẰM NGAY SAU TÊN CHƯƠNG TRÌNH", True),
    ("· content — NỘI DUNG CHƯƠNG TRÌNH: ưu đãi gì, điều kiện, đối tượng, kênh truyền thông, quà tặng. Hiện ở khối 'Chi tiết' trên M7.2.", False),
    ("· hypothesis — GIẢ THUYẾT, viết TRƯỚC khi chạy: phương án nào, tăng bao nhiêu. Hiện ở khối 'Chi tiết' trên M7.2.", False),
    ("· objective — MỤC TIÊU: TC (tăng số lượt khách) · AOV (tăng chi tiêu mỗi bill) · BRANDING.", False),
    ("· lever_primary — PHƯƠNG ÁN thuộc mục tiêu (theo TC_AOV_FnB_Marketing.pdf):", False),
    ("     TC  → Customer Base · Frequently · Party Size · Secondary Order · Sales Channel · Day-part Sales", False),
    ("     AOV → Menu Item Sold · Menu Item Value · Pricing", False),
    ("     BRANDING → Awareness · Engagement · Review          (mã + giải thích: sheet DANH_MUC)", False),
    ("", False),
    ("MÀU Ô", True),
    ("[need] CAM — cần bạn nhập.", False),
    ("[guess] XANH — máy điền, cần kiểm: ngày & cửa hàng = lần đầu/cuối thấy trên POS; mục tiêu · phương án · cơ chế đoán theo tên.", False),
    ("Ô trắng — dữ liệu gốc (tên POS, tên + mã kế hoạch). Không sửa name_pos.", False),
    ("", False),
    ("KHÔNG CẦN KHAI", True),
    ("· Nhịp chạy (một đợt / lặp theo thứ / liên tục) — hệ thống tự suy từ các ngày có hoá đơn trên POS.", False),
    ("· Giảm giá, phiếu giảm giá, số hoá đơn, doanh thu — tự lấy từ POS theo name_pos + cửa hàng + kỳ chạy.", False),
    ("· Dòng có pre_id: target + chi phí kế hoạch TỰ LẤY từ Pre-Analysis. Sửa kế hoạch thì sửa ở file Pre-Analysis.", False),
    ("", False),
    ("SAU KHI NHẬP", True),
    ("· Nháy đúp CAP_NHAT.bat → M7 · M7.1 · M7.2 cập nhật cùng lúc. Việc còn thiếu hiện ở bảng 'Việc cần khai báo' trên M7.2.", False),
    ("· Tháng sau có CTKM mới trên POS: chạy  python tools/campaign_seed.py --merge  — chỉ thêm dòng mới, không đụng ô đã nhập.", False),
]


def merge(path, dims, tgts, costs, fills):
    """Thêm dòng mới vào file đã có: bỏ chương trình trùng pre_id hoặc có tên POS đã khai."""
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill
    have = E._sheet(path, "dim_campaign")
    have_pre = {str(r.get("pre_id") or "").strip().upper() for r in have} - {""}
    have_pos = {E.norm(n) for r in have for n in str(r.get("name_pos") or "").split("|") if n.strip()}
    have_id = {str(r.get("campaign_id")) for r in have}
    keep = [i for i, d in enumerate(dims)
            if not (d.get("pre_id") and d["pre_id"] in have_pre)
            and not ({E.norm(n) for n in str(d.get("name_pos") or "").split("|") if n.strip()} & have_pos)
            and d["campaign_id"] not in have_id]
    if not keep:
        log("  không có chương trình mới — file giữ nguyên")
        return 0
    wb = load_workbook(path)
    PF = {k: PatternFill("solid", fgColor=v) for k, v in TPL.FILL.items()}
    ws = wb["dim_campaign"]
    hdr = [c.value for c in ws[1]]
    new_ids = set()
    for i in keep:
        d = dims[i]
        new_ids.add(d["campaign_id"])
        ws.append([d.get(h) for h in hdr])
        for j, h in enumerate(hdr, 1):
            k = fills.get(("dim_campaign", i, h))
            if k:
                ws.cell(ws.max_row, j).fill = PF[k]
    for sheet, rows in (("campaign_target", tgts), ("campaign_cost", costs)):
        ws = wb[sheet]
        hdr = [c.value for c in ws[1]]
        for i, t in enumerate(rows):
            if t["campaign_id"] in new_ids:
                ws.append([t.get(h) for h in hdr])
                for j, h in enumerate(hdr, 1):
                    k = fills.get((sheet, i, h))
                    if k:
                        ws.cell(ws.max_row, j).fill = PF[k]
    wb.save(path)
    log(f"  + {len(keep)} chương trình mới: " + ", ".join(dims[i]["campaign_id"] for i in keep))
    return len(keep)


# ─────────────── món LTO (chương trình chạy theo món) ───────────────
def lto_reference(item_rows):
    """Danh sách mọi món nhóm LTO đã bán trên POS — để team tra mã món khi nối chương trình."""
    E.load_facts() if not E.LTO_LINES else None
    link = {}
    for r in item_rows:
        if r.get("item_code"):
            link[E.norm(r["item_code"]).removesuffix("mk")] = r["campaign_id"]
    agg = {}
    for r in E.LTO_LINES:
        a = agg.setdefault(r["base"], dict(item_base=str(r.get("item_base") or "").upper(), item_name=r.get("item_name"),
                                          group=r.get("group"), stores=set(), first=r["date"], last=r["date"],
                                          qty=0.0, bills=set(), line_rev=0.0))
        a["stores"].add(r["store"])
        a["first"], a["last"] = min(a["first"], r["date"]), max(a["last"], r["date"])
        a["qty"] += E.to_num(r.get("qty"), 0) or 0
        a["bills"].add((r["store"], r["bill"]))
        a["line_rev"] += E.to_num(r.get("line_rev"), 0) or 0
    out = []
    for k, a in sorted(agg.items(), key=lambda kv: (str(kv[1]["group"]), -kv[1]["qty"])):
        out.append(dict(a, stores="|".join(sorted(a["stores"])), first=str(a["first"]), last=str(a["last"]),
                        qty=round(a["qty"]), bills=len(a["bills"]), line_rev=round(a["line_rev"]),
                        campaign_id=link.get(k)))
    return out


def lto_items(dims, rep):
    """campaign_item dựng sẵn: LTO Summer Crush theo món trong báo cáo S17 (tra mã trên POS);
    chương trình LTO trong kế hoạch (pre L#) → một dòng trống tô CAM để team điền mã món khi ra mắt."""
    rows, fills = [], {}
    by_name = defaultdict(set)
    for r in E.LTO_LINES:
        by_name[E.norm(r.get("item_name"))].add(str(r.get("item_base") or "").upper())
    for d in dims:
        if d.get("source") == "REPORT" and rep and rep.get("items"):
            for nm in rep["items"]:
                for code in sorted(by_name.get(E.norm(nm), [])) or [None]:
                    rows.append(dict(campaign_id=d["campaign_id"], item_code=code, item_name=nm,
                                     note=f"Món LTO theo báo cáo {rep['file']}"))
                    if not code:
                        fills[("campaign_item", len(rows) - 1, "item_code")] = "need"
        elif d.get("pre_id", "") and str(d["pre_id"]).startswith("L"):
            rows.append(dict(campaign_id=d["campaign_id"], item_code=None, item_name=None,
                             note=f"LTO trong kế hoạch {d['pre_id']} — điền mã món khi ra mắt (tra sheet DS_MON_LTO)"))
            fills[("campaign_item", len(rows) - 1, "item_code")] = "need"
    return rows, fills


def add_item_sheets(path, dims, rep):
    """File đã có (có thể đã nhập tay): CHỈ thêm/làm mới sheet campaign_item (nếu chưa có) và DS_MON_LTO."""
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = load_workbook(path)
    HDR = PatternFill("solid", fgColor="1F1F26")
    PF = {k: PatternFill("solid", fgColor=v) for k, v in TPL.FILL.items()}

    def new_sheet(name, cols, rows, fills):
        ws = wb.create_sheet(name)
        for j, (col, desc) in enumerate(cols, 1):
            c = ws.cell(1, j, col)
            c.font = Font(bold=True, color="F3F2EE")
            c.fill = HDR
            c.alignment = Alignment(horizontal="center")
            c.comment = Comment(desc, "NOIRE")
            ws.column_dimensions[c.column_letter].width = 18 if col != "note" else 60
        for i, r in enumerate(rows):
            ws.append([r.get(col) for col, _ in cols])
            for j, (col, _) in enumerate(cols, 1):
                k = fills.get((name, i, col))
                if k:
                    ws.cell(ws.max_row, j).fill = PF[k]
        ws.freeze_panes = "B2"
        return ws

    added = []
    if "campaign_item" not in wb.sheetnames:
        rows, fills = lto_items(dims, rep)
        new_sheet("campaign_item", TPL.ITEM_COLS, rows, fills)
        added.append(f"campaign_item ({len(rows)} dòng)")
    have = E._sheet(path, "campaign_item") if "campaign_item" not in [a.split(" ")[0] for a in added] else rows
    if "DS_MON_LTO" in wb.sheetnames:
        del wb["DS_MON_LTO"]
    ref = lto_reference(have)
    new_sheet("DS_MON_LTO", TPL.LTO_REF_COLS, ref, {})
    added.append(f"DS_MON_LTO ({len(ref)} món)")
    wb.save(path)
    log("  + " + " · ".join(added))


def main(argv):
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    d = l0_dir("S23_campaign")
    if not d:
        log("  chưa có thư mục 07_Campaign_Tracking — chạy python tools/l0_setup.py trước")
        return 1
    path = os.path.join(d, OUT_NAME)
    plan, _ = E.load_plan()
    daily, promo, ads, first, last = E.load_facts()
    groups = pos_groups(promo, last)
    rep = report_lto()
    dims, tgts, costs, fills = build(plan, groups, rep, last)

    n_pre_pos = sum(1 for x in dims if x.get("source") == "PRE_POS")
    n_need = sum(1 for (s, _, _), v in fills.items() if v == "need")
    log(f"  kế hoạch {len(plan)} CT · POS {len(groups)} nhóm tên (bỏ ưu đãi nội bộ) · báo cáo LTO {'có' if rep else 'không'}")
    log(f"  → {len(dims)} chương trình: {n_pre_pos} khớp kế hoạch↔POS · "
        f"{sum(1 for x in dims if x.get('source') == 'PRE')} chỉ có kế hoạch · "
        f"{sum(1 for x in dims if x.get('source') in ('POS', 'REPORT'))} chỉ có trên POS · {n_need} ô cần nhập")
    if "--dry" in argv:
        for x in dims:
            log(f"   {x['campaign_id']:<34} {x.get('source'):<7} {x.get('pre_id') or '':<4} {x['name'][:40]:<40} | {str(x.get('name_pos') or '')[:60]}")
        return 0
    if os.path.exists(path):
        if "--merge" in argv:
            merge(path, dims, tgts, costs, fills)
            add_item_sheets(path, E._sheet(path, "dim_campaign"), rep)
            return 0
        log(f"  {OUT_NAME} đã có — KHÔNG ghi đè dữ liệu bạn đã nhập. Dùng --merge để chỉ thêm chương trình mới.")
        return 1
    items, ifills = lto_items(dims, rep)
    fills.update(ifills)
    TPL.write_book(path, GUIDE, dims, tgts, costs, fills=fills, items=items, lto_ref=lto_reference(items))
    log(f"  đã ghi {os.path.relpath(path)}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
