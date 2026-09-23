# -*- coding: utf-8 -*-
"""
M7.1 · PRE-ANALYTICS — BỘ MÁY ĐÁNH GIÁ CHƯƠNG TRÌNH TRƯỚC KHI CHẠY
==================================================================
    python tools/preeval.py

Đọc  L0_input/03_MARKETING/05_Promotion_Ke_Hoach/01_So_Danh_Gia/Pre_Analysis_*.xlsx  (S24)
   + data_input/monthly/*.xlsx   daily (TC · khách · doanh thu theo cửa hàng × ngày)
                                  fact_promo_day (giảm giá · hệ số thuế/phí theo cửa hàng)
   + data_input/02_snapshot.xlsx  product (giá bán · giá vốn BOM từng món)
Ghi  data_input/04_preeval.xlsx   pre_eval · pre_eval_scheme · pre_eval_fin · pre_eval_base

KHUNG PHIẾU = PP672 × SALES = TC × AOV. Mọi số tiền tính trên GIÁ MENU (trước phí DV + VAT, như POS
"Thanh toán trước giảm giá") rồi quy ra:
    Tổng tiền (gồm VAT)      = (menu − giảm) × hệ số thuế/phí của chính các cửa hàng đó (đo từ POS)
    Doanh thu thuần (−VAT)   = Tổng tiền ÷ 1,08

1) NỀN (Base) — kỳ nền 56 ngày trước ngày bắt đầu (bỏ ngày lễ; lọc đúng các thứ chạy), từng cửa hàng,
   quy về số ngày chương trình: Base = run-rate/ngày mở bán × số ngày chạy.
2) KINH TẾ 1 HOÁ ĐƠN THAM GIA (từng scheme):
    giá trị hoá đơn  = max(min_bill, AOV nền + món chương trình × %gọi thêm)   (GROUP_SIZE: TA × số khách)
    giảm giá         = theo ưu đãi (PCT_OFF_BILL · PCT_OFF_ITEMS · FIXED_OFF · FIXED_PRICE · GIFT_ITEM)
    giá vốn          = món REQUIRED + món GIFT (BOM hoặc nhập) + món khác × COGS% nền
    quà vật phẩm     = đơn giá × (1 + VAT)
3) SỐ HOÁ ĐƠN THAM GIA = est_bills | participation% × TC nền, không vượt stock_qty.
4) FINANCIAL EVALUATION:
    Có KM            = Σ hoá đơn tham gia × kinh tế 1 hoá đơn
    Bị ăn mòn (cannib)= %cannib × hoá đơn tham gia × giá trị 1 hoá đơn NỀN (khách vốn sẽ đến)
    Không KM         = Base − Bị ăn mòn
    Tổng             = Không KM + Có KM          Tăng thêm = Tổng − Base
    %Cannib (PP672)  = (Base − Không KM) ÷ Không KM
5) EBITDA TĂNG THÊM = LN gộp tăng thêm − quà vật phẩm − chi phí chương trình − chi phí vận hành biến đổi
   (ty_le_chi_phi × doanh thu thuần tăng thêm).
6) 3 KỊCH BẢN ($preeval.scenarios) + điểm hoà vốn (hoá đơn cần có · %cannib tối đa) + quyết định.
"""
from __future__ import annotations

import glob
import os
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from monthly_lib import (  # noqa: E402
    BRAND_OF_STORE, CONTRACT, DATA_INPUT, MONTHLY_DIR, STORE_META, date_of, l0_files, norm,
    read_workbook, to_num, write_workbook,
)

P = CONTRACT["$preeval"]
CAMP = CONTRACT["$campaign"]
VAT = P["vat"]
OUT = os.path.join(DATA_INPUT, "04_preeval.xlsx")
SID = "S24_preeval"
LEVER_GROUP = {x["code"]: x.get("group") for x in CAMP["levers"]}

EXCLUDE = set()
for _a, _b, _ in CAMP.get("base_exclude", []):
    _x, _y = date.fromisoformat(_a), date.fromisoformat(_b)
    while _x <= _y:
        EXCLUDE.add(_x)
        _x += timedelta(1)


def log(*a):
    print(*a, flush=True)


def _d(v):
    s = date_of(v)
    return date.fromisoformat(s) if s else None


def _f(v, default=None):
    x = to_num(v)
    return default if x is None else x


def _div(a, b):
    return a / b if b else None


# ─────────────────────────── đọc sổ ───────────────────────────
def read_sheet(path, name):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if name not in wb.sheetnames:
            return []
        rows = list(wb[name].iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    hdr = [str(c).strip() if c is not None else "" for c in rows[0]]
    out = []
    for r in rows[1:]:
        if not r or r[0] is None or str(r[0]).strip() == "":
            continue
        rec = {}
        for j, h in enumerate(hdr):
            if not h:
                continue
            v = r[j] if j < len(r) else None
            if isinstance(v, (datetime, date)):
                v = v.strftime("%Y-%m-%d")
            elif isinstance(v, str):
                v = v.strip() or None
            rec[h] = v
        out.append(rec)
    return out


# ─────────────────────────── dữ liệu nền ───────────────────────────
def load_facts():
    daily = {}
    promo = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0])          # (store, date) → net, gross, disc, bills
    for f in sorted(glob.glob(os.path.join(MONTHLY_DIR, "*.xlsx"))):
        wb = read_workbook(f, {"daily", "fact_promo_day"})
        for r in wb.get("daily", []):
            d = _d(r.get("date"))
            if d:
                daily[(r["store"], d)] = (_f(r.get("net"), 0), _f(r.get("tc"), 0), _f(r.get("guest"), 0))
        for r in wb.get("fact_promo_day", []):
            d = _d(r.get("date"))
            if d and _f(r.get("bills"), 0):
                x = promo[(r["store"], d)]
                x[0] += _f(r.get("net"), 0)
                x[1] += _f(r.get("gross"), 0)
                x[2] += _f(r.get("disc"), 0)
                x[3] += _f(r.get("bills"), 0)
    days = sorted({d for _, d in daily})
    snap = read_workbook(os.path.join(DATA_INPUT, "02_snapshot.xlsx"), {"product"})
    menu = {}
    for r in snap.get("product", []):
        q = _f(r.get("qty"), 0)
        if not q:
            continue
        cg = _f(r.get("cogs"), 0)
        menu[norm(r.get("ma"))] = dict(name=r.get("name"), price=_f(r.get("rev"), 0) / q,
                                       cogs=(cg / q) if cg else None)
    return daily, promo, menu, (days[0] if days else None), (days[-1] if days else None)


def item_ref(menu, code):
    k = norm(code)
    return menu.get(k) or menu.get(k.removesuffix("mk")) or menu.get(k + "mk")


def base_window(d0, dows, first, last, bf=None, bt=None):
    """Ngày của kỳ nền: tự chọn (base_from/base_to) hoặc 56 ngày ngay trước ngày bắt đầu (bỏ lễ).
    Chương trình ở tương lai xa hơn dữ liệu → 56 ngày cuối có số."""
    if bf and bt:
        ds = [bf + timedelta(i) for i in range((bt - bf).days + 1)]
        note = "kỳ nền tự chọn"
    else:
        end = min(d0 - timedelta(1), last) if last else d0 - timedelta(1)
        ds, d = [], end
        while len(ds) < P["base_days"] and first and d >= first:
            if d not in EXCLUDE:
                ds.append(d)
            d -= timedelta(1)
        note = ("56 ngày trước ngày bắt đầu" if end == d0 - timedelta(1)
                else "chương trình sau ngày có số cuối — lấy 56 ngày gần nhất")
    if dows:
        ds = [d for d in ds if d.weekday() in dows]
    return sorted(ds), note


QUICK_COLS = ("condition", "min_bill", "min_guests", "benefit", "benefit_value", "discount_cap",
              "merch_name", "merch_unit_cost", "merch_vat_pct", "stock_qty", "bill_value", "item_add_pct")


def _codes(text, role, pid):
    """'LTD0009*0.667|FCK0005' → dòng mon. Thiếu '*SL' = 1."""
    out = []
    for part in str(text or "").replace(",", "|").split("|"):
        part = part.strip()
        if not part:
            continue
        code, _, q = part.partition("*")
        out.append(dict(program_id=pid, scheme_id="S1", role=role, item_code=code.strip(),
                        qty=to_num(q.strip()) if q.strip() else 1))
    return out


def quick_scheme(prog):
    """Chương trình KHÔNG có dòng ở co_che → dựng 1 scheme từ cột cơ chế nhanh trên chuong_trinh.
    Để một ý tưởng chỉ cần MỘT dòng; chương trình nhiều scheme mới cần sheet co_che."""
    pid = prog["program_id"]
    sch = dict(program_id=pid, scheme_id="S1", scheme_name=prog.get("offer") or prog.get("name"))
    for k in QUICK_COLS:
        sch[k] = prog.get(k)
    items = _codes(prog.get("item_codes"), "REQUIRED", pid) + _codes(prog.get("gift_codes"), "GIFT", pid)
    return sch, items


# ─────────────────────────── đánh giá một chương trình ───────────────────────────
def evaluate(prog, schemes, items, costs, opex, daily, promo, menu, first, last):
    pid = prog["program_id"]
    mode = "CHI_TIET"
    if not schemes:
        q, qi = quick_scheme(prog)
        schemes, items, mode = [q], qi, "NHANH"
    season = _f(prog.get("season_factor"), 1.0) or 1.0
    brand = (prog.get("brand") or "ALL").upper()
    stores = [s.strip().upper() for s in str(prog.get("store_scope") or "").split("|") if s.strip()]
    d0, d1 = _d(prog.get("date_from")), _d(prog.get("date_to"))
    dows = {int(x) for x in str(prog.get("dow") or "").replace(",", "|").split("|") if str(x).strip().isdigit()}
    issues = []
    if not stores or not d0 or not d1 or d1 < d0:
        return None, [], [], [], [f"{pid}: thiếu cửa hàng hoặc ngày chạy"]
    pdays = [d0 + timedelta(i) for i in range((d1 - d0).days + 1)]
    if dows:
        pdays = [d for d in pdays if d.weekday() in dows]
    n_p = len(pdays)
    bw, bnote = base_window(d0, dows, first, last, _d(prog.get("base_from")), _d(prog.get("base_to")))

    # ── 1. NỀN theo cửa hàng ──
    base_rows, B = [], dict(net=0.0, tc=0.0, guest=0.0, gross=0.0, disc=0.0)
    for st in stores:
        vals = [daily[(st, d)] for d in bw if (st, d) in daily]
        open_days = len(vals)
        if not open_days:
            issues.append(f"{pid}: {st} không có số trong kỳ nền")
            base_rows.append(dict(program_id=pid, store=st, base_days=0, note="không có số trong kỳ nền"))
            continue
        net, tc, gu = (sum(v[i] for v in vals) for i in range(3))
        pr = [promo[(st, d)] for d in bw if (st, d) in promo]
        pnet, pgross, pdisc, pbills = (sum(x[i] for x in pr) for i in range(4))
        tf = pnet / (pgross - pdisc) if pbills >= 20 and pgross > pdisc else VAT
        tf = tf if 1.0 <= tf <= 1.25 else VAT
        gross = net / tf + pdisc
        k = n_p / open_days * season          # hệ số mùa vụ: kỳ nền khác mùa với kỳ chạy (vd Q4 lễ cuối năm)
        B["net"] += net * k
        B["tc"] += tc * k
        B["guest"] += gu * k
        B["gross"] += gross * k
        B["disc"] += pdisc * k
        base_rows.append(dict(program_id=pid, store=st, base_from=str(bw[0]), base_to=str(bw[-1]), base_days=open_days,
                              net_incl=round(net), tc=round(tc), guests=round(gu),
                              aov_incl=round(net / tc) if tc else None, ta_incl=round(net / gu) if gu else None,
                              tc_day=round(tc / open_days, 1), net_day=round(net / open_days),
                              tax_factor=round(tf, 4), disc_share=round(pdisc / gross, 4) if gross else None,
                              note=bnote + (f" · lọc thứ {sorted(dows)}" if dows else "")))
    if not B["tc"]:
        return None, [], [], base_rows, issues + [f"{pid}: không có TC nền"]
    TF = B["net"] / (B["gross"] - B["disc"]) if B["gross"] > B["disc"] else VAT
    aov_menu = B["gross"] / B["tc"]                    # giá trị menu 1 hoá đơn nền
    ta_menu = B["gross"] / B["guest"] if B["guest"] else aov_menu
    disc_bill = B["disc"] / B["tc"]
    ocogs = _f(prog.get("other_cogs_pct")) or P["base_cogs_pct"].get(brand, P["base_cogs_pct"]["ALL"])

    # ── 2. KINH TẾ 1 HOÁ ĐƠN theo scheme ──
    lever = (prog.get("lever_primary") or "").upper()
    cannib0 = _f(prog.get("cannib_pct"))
    cannib_src = "nhập ở sổ"
    if cannib0 is None:
        cannib0 = P["cannib_default"].get(lever, P["cannib_default"]["_default"])
        cannib_src = f"mặc định theo phương án {lever or '(chưa chọn)'}"
    part = _f(prog.get("participation_pct"))
    if _f(prog.get("est_bills")):
        demand, part_src = _f(prog.get("est_bills")), "est_bills nhập ở sổ"
    elif part is not None:
        demand, part_src = part * B["tc"], f"{part:.1%} TC nền (nhập ở sổ)"
    else:
        demand, part_src = P["participation_default"] * B["tc"], f"{P['participation_default']:.0%} TC nền (mặc định)"
    by_scheme = defaultdict(list)
    for it in items:
        by_scheme[str(it.get("scheme_id"))].append(it)
    n_s = len(schemes) or 1
    econ = []
    for s in schemes:
        sid = str(s.get("scheme_id"))
        notes = []
        req_v = req_c = gift_v = gift_c = 0.0
        for it in by_scheme.get(sid, []):
            q = _f(it.get("qty"), 1)
            ref = item_ref(menu, it.get("item_code")) or {}
            price = _f(it.get("price")) or ref.get("price")
            cogs = _f(it.get("unit_cogs")) or ref.get("cogs")
            if price is None:
                notes.append(f"{it.get('item_code')}: chưa có giá — tính 0")
                price = 0.0
            if cogs is None:
                cogs = price * ocogs
                notes.append(f"{it.get('item_name') or it.get('item_code')}: chưa có giá vốn BOM — tạm {ocogs:.0%}")
            if str(it.get("role") or "REQUIRED").upper() == "GIFT":
                gift_v += q * price
                gift_c += q * cogs
            else:
                req_v += q * price
                req_c += q * cogs
        benefit = (s.get("benefit") or "NONE").upper()
        cond = (s.get("condition") or "NONE").upper()
        ref_v = ta_menu * _f(s.get("min_guests"), 1) if cond == "GROUP_SIZE" else aov_menu
        add = _f(s.get("item_add_pct"))
        add = P["item_add_pct"].get(benefit, 0.0) if add is None else add
        bill_v = _f(s.get("bill_value")) or max(_f(s.get("min_bill"), 0), ref_v + req_v * add, req_v)
        val = _f(s.get("benefit_value"), 0)
        if benefit == "PCT_OFF_BILL":
            disc = bill_v * val
        elif benefit == "PCT_OFF_ITEMS":
            disc = req_v * val
        elif benefit == "FIXED_OFF":
            disc = min(val, bill_v)
        elif benefit == "FIXED_PRICE":
            disc = max(0.0, req_v - val)
        elif benefit == "GIFT_ITEM":
            disc = gift_v
        else:
            disc = 0.0
        cap = _f(s.get("discount_cap"))
        if cap is not None and benefit != "GIFT_ITEM":
            disc = min(disc, cap)
        other_v = max(0.0, bill_v - req_v)
        merch = _f(s.get("merch_unit_cost"), 0) * (1 + _f(s.get("merch_vat_pct"), 0.08)) if benefit == "GIFT_MERCH" else 0.0
        share = _f(s.get("share_pct"))
        econ.append(dict(
            s=s, sid=sid, benefit=benefit, cond=cond, req_v=req_v, req_c=req_c, gift_v=gift_v, gift_c=gift_c,
            bill_v=bill_v, disc=disc, other_v=other_v, merch=merch, share=share if share is not None else 1 / n_s,
            est=_f(s.get("est_bills")), stock=_f(s.get("stock_qty")),
            note=" · ".join(notes + ["giá trị hoá đơn = max(hoá đơn tối thiểu; %s %s + món chương trình × %.0f%%)"
                                     % ("TA × số khách" if cond == "GROUP_SIZE" else "AOV nền",
                                        f"{ref_v:,.0f}".replace(",", "."), add * 100)])))

    fixed_cost = sum(_f(c.get("amount"), 0) * (1 + _f(c.get("vat_pct"), 0.08)) for c in costs)
    var_opex = sum(_f(o.get(brand if brand in ("NCB", "NDC", "NJFB") else "NCB"), 0)
                   for o in opex if str(o.get("type")).upper() == "VARIABLE" and int(_f(o.get("act_on_incr"), 0)) == 1)

    def run(bills_mult=1.0, cannib_add=0.0, cogs_add=0.0, cannib_fix=None, bills_scale=None):
        oc = ocogs + cogs_add
        c = cannib_fix if cannib_fix is not None else min(1.0, max(0.0, cannib0 + cannib_add))
        tot_bills = gross_p = disc_p = cogs_p = merch_t = 0.0
        stock = 0.0
        per = []
        for e in econ:
            b = e["est"] if e["est"] is not None else demand * e["share"]
            b *= bills_mult
            if bills_scale is not None:
                b = bills_scale * e["share"]
            if e["stock"]:
                stock += e["stock"]
                b = min(b, e["stock"])
            cogs_bill = e["req_c"] + e["gift_c"] + e["other_v"] * oc
            tot_bills += b
            gross_p += b * (e["bill_v"] + e["gift_v"])
            disc_p += b * e["disc"]
            cogs_p += b * cogs_bill
            merch_t += b * e["merch"]
            per.append((e, b, cogs_bill))
        net = lambda g, dsc: (g - dsc) * TF / VAT          # noqa: E731
        base = dict(tc=B["tc"], gross=B["gross"], disc=B["disc"])
        base["net"] = net(base["gross"], base["disc"])
        base["cogs"] = base["net"] * oc
        cb = c * tot_bills
        cann = dict(tc=cb, gross=cb * aov_menu, disc=cb * disc_bill)
        cann["net"] = net(cann["gross"], cann["disc"])
        cann["cogs"] = cann["net"] * oc
        withp = dict(tc=tot_bills, gross=gross_p, disc=disc_p, net=net(gross_p, disc_p), cogs=cogs_p)
        without = {k: base[k] - cann[k] for k in base}
        total = {k: without[k] + withp[k] for k in base}
        incr = {k: total[k] - base[k] for k in base}
        gp_incr = incr["net"] - incr["cogs"]
        opex_incr = max(0.0, incr["net"]) * var_opex
        ebitda = gp_incr - merch_t - fixed_cost - opex_incr
        spend = disc_p * TF / VAT + merch_t + fixed_cost
        return dict(c=c, bills=tot_bills, stock=stock, base=base, cann=cann, withp=withp, without=without,
                    total=total, incr=incr, gp_incr=gp_incr, opex=opex_incr, merch=merch_t, ebitda=ebitda,
                    spend=spend, per=per, oc=oc)

    results, fin, scheme_rows = [], [], []
    for sc in P["scenarios"]:
        R = run(sc["bills_mult"], sc["cannib_add"], sc["cogs_add"])
        # điểm hoà vốn theo số hoá đơn (tuyến tính, bỏ trần stock) và theo %cannib
        e0 = -fixed_cost
        a = (R["ebitda"] - e0) / R["bills"] if R["bills"] else 0
        be_bills = (fixed_cost / a if fixed_cost else 0.0) if a > 0 else None
        E0 = run(sc["bills_mult"], cogs_add=sc["cogs_add"], cannib_fix=0.0)["ebitda"]
        E1 = run(sc["bills_mult"], cogs_add=sc["cogs_add"], cannib_fix=1.0)["ebitda"]
        max_c = (E0 / (E0 - E1)) if (E0 > 0 and E1 < 0) else (1.0 if E1 >= 0 else 0.0)
        W = R["withp"]
        flags = []
        if W["net"] and W["cogs"] / W["net"] > P["gates"]["max_cogs_pct"]:
            flags.append("COGS hoá đơn tham gia %.0f%% > %.0f%%" % (W["cogs"] / W["net"] * 100, P["gates"]["max_cogs_pct"] * 100))
        if W["net"] and R["spend"] / W["net"] > P["gates"]["max_promo_cost_pct_net"]:
            flags.append("chi phí KM %.0f%% doanh thu thuần > %.0f%% — cần bằng chứng incremental (A/B, đối chứng)"
                         % (R["spend"] / W["net"] * 100, P["gates"]["max_promo_cost_pct_net"] * 100))
        results.append(dict(
            program_id=pid, scenario=sc["code"], name=prog.get("name"), brand=brand, stores="|".join(stores),
            date_from=str(d0), date_to=str(d1), days=n_p, objective=(prog.get("objective") or LEVER_GROUP.get(lever)),
            lever=lever or None, status=prog.get("status"), bills=round(R["bills"]),
            bills_incr=round(R["bills"] * (1 - R["c"])), cannib_pct=round(R["c"], 4),
            rev_incl=round(W["net"] * VAT), net_incr=round(R["incr"]["net"]), gp_incr=round(R["gp_incr"]),
            promo_cost=round(R["spend"]), program_cost=round(fixed_cost + R["merch"]), opex_incr=round(R["opex"]),
            ebitda_incr=round(R["ebitda"]), ebitda_pct=_div(R["ebitda"], W["net"]),
            roi=_div(R["ebitda"], R["spend"]), breakeven_bills=round(be_bills) if be_bills is not None else None,
            max_cannib=round(max_c, 4), redemption_needed=_div(R["stock"], B["tc"]) if R["stock"] else None,
            quarter=f"{d0.year}-Q{(d0.month - 1) // 3 + 1}", season_factor=season, scheme_mode=mode,
            tc_base=round(B["tc"]), tc_share=_div(R["bills"], B["tc"]), participation_src=part_src,
            cannib_src=cannib_src, other_cogs_pct=R["oc"], opex_pct=var_opex, base_note=bnote,
            safety_bills=_div(R["bills"], be_bills) if be_bills else None,
            stock_days=round(R["stock"] / (R["bills"] / n_p), 1) if R["stock"] and R["bills"] and n_p else None,
            gate_flags=" · ".join(flags) or None, campaign_id=prog.get("campaign_id")))
        for key, label in (("tc", "Hoá đơn (TC)"), ("gross", "Gross sales — gồm VAT"), ("disc", "Giảm giá — gồm VAT"),
                           ("net", "Doanh thu thuần — chưa VAT"), ("cogs", "Giá vốn (COGS)"), ("gp", "Lợi nhuận gộp")):
            def v(col):
                x = R[col]
                if key == "gp":
                    return x["net"] - x["cogs"]
                return x[key] * (TF if key in ("gross", "disc") else 1)
            b_, w_, p_, t_ = v("base"), v("without"), v("withp"), v("total")
            fin.append(dict(program_id=pid, scenario=sc["code"], row=key, label=label, base=round(b_),
                            without=round(w_), with_promo=round(p_), total=round(t_),
                            cannib_pct=_div(b_ - w_, w_) if key in ("net", "tc") else None,
                            incr=round(t_ - b_), incr_pct=_div(t_ - b_, b_)))
        if sc["code"] == "CO_SO":
            for e, b, cogs_bill in R["per"]:
                rev_after = (e["bill_v"] - e["disc"]) * TF
                net_after = (e["bill_v"] - e["disc"]) * TF / VAT
                scheme_rows.append(dict(
                    program_id=pid, scheme_id=e["sid"], scheme_name=e["s"].get("scheme_name"), condition=e["cond"],
                    benefit=e["benefit"], bills=round(b), bill_value=round((e["bill_v"] + e["gift_v"]) * TF),
                    discount=round(e["disc"] * TF), rev_after_disc=round(rev_after), ta=round(rev_after),
                    cogs=round(cogs_bill), cogs_pct=_div(cogs_bill, net_after), margin_pct=(1 - cogs_bill / net_after) if net_after else None,
                    merch_cost=round(e["merch"]), promo_cost=round(e["disc"] * TF / VAT + e["merch"]), basis_note=e["note"]))

    base_r = {r["scenario"]: r for r in results}
    eb, ec = base_r["CO_SO"]["ebitda_incr"], base_r["THAN_TRONG"]["ebitda_incr"]
    obj = (base_r["CO_SO"]["objective"] or "").upper()
    if obj == "BRANDING":
        dec, why = "BRANDING", "mục tiêu branding — duyệt theo ngân sách và KPI tiếp cận"
    elif eb > 0 and ec >= 0:
        dec, why = "DUYET", "có lãi ở cả kịch bản Cơ sở và Thận trọng"
    elif eb > 0:
        dec, why = "CHAY_THU", ("có lãi ở Cơ sở nhưng lỗ %s ở Thận trọng — chạy thử 1–2 tuần, theo dõi số hoá đơn và %%cannib"
                                % f"{-ec:,.0f}".replace(",", "."))
    else:
        dec, why = "SUA_CO_CHE", "lỗ ở kịch bản Cơ sở — đổi ưu đãi / ngưỡng / chi phí rồi đánh giá lại"
    for r in results:
        r["decision"], r["decision_note"] = dec, why
    return results, fin, scheme_rows, base_rows, issues


# ─────────────────────────── đầu vào chuẩn: sổ + file deck quý ───────────────────────────
FIELDS = P["input_fields"]["fields"]


def _empty(v):
    return v is None or (isinstance(v, str) and not v.strip())


def merge_inputs(progs, costs):
    """Sổ Pre_Analysis (S24) + file deck quý (S16, Q4/2026+) → MỘT danh sách chương trình theo mẫu chuẩn.
    File deck đã được tools/pre_analysis.py chuyển sang đúng cột `chuong_trinh`. Dòng sổ có cùng
    program_id bổ sung / ghi đè TỪNG Ô: ô sổ có số thì dùng số sổ, ô sổ trống thì giữ số deck.
    Chi phí: sổ có dòng chi_phi cho chương trình → dùng sổ; chưa có → chi phí MKT của deck."""
    import pre_analysis
    deck = pre_analysis.deck_programs()
    so = {p["program_id"]: p for p in progs}
    so_cost = {c.get("program_id") for c in costs}
    out, extra_costs = [], []
    for d in deck["programs"]:
        s = so.pop(d["program_id"], None) or {}
        m = {k: v for k, v in d.items() if not k.startswith("_")}
        src = {k: "DECK" for k, v in m.items() if not _empty(v)}
        for k, v in s.items():
            if not _empty(v):
                m[k], src[k] = v, "SO"
        m.update(_src=src, _hint=d["_hint"], _quarter=d["_quarter"],
                 _source=f"DECK: {d['_file']}" + (" + sổ" if s else ""))
        out.append(m)
        if d["program_id"] not in so_cost:
            extra_costs += [dict(c, _src="DECK") for c in deck["costs"] if c["program_id"] == d["program_id"]]
    for p in progs:                                   # chương trình chỉ có ở sổ
        if p["program_id"] in so:
            out.append(dict(p, _src={k: "SO" for k, v in p.items() if not _empty(v)}, _hint={}, _source="SO"))
    return out, costs + extra_costs, deck["issues"]


def missing_fields(prog, has_scheme):
    """Trường bắt buộc còn trống ($preeval.input_fields) → danh sách mã trường."""
    branding = str(prog.get("objective") or "").upper() == "BRANDING"
    miss = []
    for f in FIELDS:
        if f["level"] == "required" and _empty(prog.get(f["code"])):
            miss.append(f["code"])
        elif f["level"] == "required_promo" and not branding and not has_scheme and _empty(prog.get(f["code"])):
            miss.append(f["code"])
    return miss


def input_rows(prog, costs, miss):
    """Dữ liệu đầu vào đã chuẩn hoá của 1 chương trình — mỗi trường một dòng (bảng pre_eval_input)."""
    out = []
    for f in FIELDS:
        k = f["code"]
        if k == "chi_phi":
            v = sum(_f(c.get("amount"), 0) for c in costs) or None
            src = ("DECK" if all(c.get("_src") == "DECK" for c in costs) else "SO") if costs else None
        else:
            v, src = prog.get(k), prog["_src"].get(k)
        out.append(dict(program_id=prog["program_id"], field=k, value=None if _empty(v) else v, source=src,
                        level=f["level"], status="THIEU" if k in miss else ("OK" if not _empty(v) else "TRONG"),
                        hint=prog["_hint"].get(k)))
    return out


def pending_row(prog, miss, why=None):
    """Chương trình CHƯA tính được (thiếu trường bắt buộc) vẫn lên M7.1 — quyết định THIEU_SO."""
    label = {f["code"]: f["label"] for f in FIELDS}
    d0 = _d(prog.get("date_from"))
    q = prog.get("_quarter") or (f"{d0.year}-Q{(d0.month - 1) // 3 + 1}" if d0 else None)
    note = why or ("Cần bổ sung: " + " · ".join(label.get(m, m) for m in miss))
    return dict(program_id=prog["program_id"], scenario="CO_SO", name=prog.get("name"),
                brand=(prog.get("brand") or "ALL").upper(),
                stores=str(prog.get("store_scope") or "").replace(" ", ""), date_from=prog.get("date_from"),
                date_to=prog.get("date_to"), objective=prog.get("objective"), lever=prog.get("lever_primary"),
                status=prog.get("status"), decision="THIEU_SO", decision_note=note, quarter=q,
                campaign_id=prog.get("campaign_id"), input_source=prog["_source"], missing="|".join(miss) or None)


def main():
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    files = l0_files(SID)
    path = max(files, key=os.path.getmtime) if files else None
    progs = read_sheet(path, "chuong_trinh") if path else []
    schemes, items, costs = ((read_sheet(path, "co_che"), read_sheet(path, "mon"), read_sheet(path, "chi_phi"))
                             if path else ([], [], []))
    opex = read_sheet(path, "ty_le_chi_phi") if path else []
    progs, costs, issues = merge_inputs(progs, costs)
    if not progs:
        log("  – chưa có sổ Pre_Analysis_*.xlsx và file deck quý — bỏ qua M7.1 đánh giá")
        write_workbook(OUT, {"pre_eval": [], "pre_eval_scheme": [], "pre_eval_fin": [], "pre_eval_base": [],
                             "pre_eval_input": []}, title="M7.1 · PRE-ANALYTICS — chưa có đầu vào")
        return 0
    daily, promo, menu, first, last = load_facts()
    out = defaultdict(list)
    for p in progs:
        pid = p["program_id"]
        sch = [s for s in schemes if s.get("program_id") == pid]
        cst = [c for c in costs if c.get("program_id") == pid]
        miss = missing_fields(p, bool(sch))
        out["pre_eval_input"] += input_rows(p, cst, miss)
        if miss:
            out["pre_eval"].append(pending_row(p, miss))
            continue
        res, fin, sc_rows, base, iss = evaluate(
            p, sch, [i for i in items if i.get("program_id") == pid], cst, opex, daily, promo, menu, first, last)
        issues += iss
        if not res:
            out["pre_eval"].append(pending_row(p, [], "Chưa tính được: " + "; ".join(iss)))
            out["pre_eval_base"] += base
            continue
        for r in res:
            r.update(input_source=p["_source"], missing=None)
            if p.get("_quarter"):                        # chương trình từ file deck thuộc quý của file đó
                r["quarter"] = p["_quarter"]
        out["pre_eval"] += res
        out["pre_eval_fin"] += fin
        out["pre_eval_scheme"] += sc_rows
        out["pre_eval_base"] += base
    write_workbook(OUT, dict(out), title="M7.1 · PRE-ANALYTICS — sinh tự động bởi tools/preeval.py")
    base = [r for r in out["pre_eval"] if r["scenario"] == "CO_SO"]
    n_deck = sum(1 for p in progs if p["_source"] != "SO")
    log(f"  sổ: {os.path.basename(path) if path else '—'} · {len(progs)} chương trình ({n_deck} từ file deck quý)")
    for r in base:
        if r["decision"] == "THIEU_SO":
            log(f"   {r['program_id']:<26} THIEU_SO    {r['decision_note']}")
        else:
            log(f"   {r['program_id']:<26} {r['decision']:<11} hoá đơn {r['bills']:>5} · tăng thêm {r['net_incr']:>13,} · EBITDA {r['ebitda_incr']:>12,}".replace(",", "."))
    for i in issues:
        log("   ! " + i)
    return 0


if __name__ == "__main__":
    sys.exit(main())
